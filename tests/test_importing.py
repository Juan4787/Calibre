import io
import warnings
import zipfile
from datetime import datetime

import pytest
import xlwt
from hypothesis import given, settings
from hypothesis import strategies as st
from openpyxl import Workbook
from pydantic import ValidationError

from freight_audit.canonical import bytes_hash
from freight_audit.fixtures import column, constant
from freight_audit.importing import (
    Cell,
    ColumnMapping,
    ImportErrorDetail,
    ImportMapping,
    convert,
    import_data,
    parse_decimal,
)


def mapping(**extra):
    return ImportMapping.model_validate(
        {
            "id": "test",
            "version": "1",
            "entity": "shipments",
            "columns": [
                column("id", "id"),
                column("ref", "reference"),
                constant("carrier", "C"),
                column("weight", "attributes.weight", "decimal", unit="kg"),
                column("date", "attributes.service_date", "date"),
            ],
            **extra,
        }
    )


def csv_data(rows):
    return ("id;ref;weight;date\n" + rows).encode("utf-8")


def xlsx_data(rows, formats=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in rows:
        sheet.append(row)
    for coordinate, fmt in (formats or {}).items():
        sheet[coordinate].number_format = fmt
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_xlsx_hidden_business_data_blocks_import_until_made_visible():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["id", "ref", "weight", "date"])
    sheet.append(["S1", "R1", 10, "30/09/2026"])
    sheet.row_dimensions[2].hidden = True
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(ImportErrorDetail, match="filas 2"):
        import_data(output.getvalue(), "hidden.xlsx", mapping())

    sheet.row_dimensions[2].hidden = False
    sheet.column_dimensions["C"].hidden = True
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(ImportErrorDetail, match="columnas 3"):
        import_data(output.getvalue(), "hidden.xlsx", mapping())

    sheet.column_dimensions["C"].hidden = False
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    result = import_data(output.getvalue(), "visible.xlsx", mapping())
    assert result["accepted"] == 1 and result["rejected"] == []
    record = result["records"][0]
    assert record["attributes"]["weight"]["value"] == "10"
    assert record["provenance"]["attributes.weight"]["row"] == 2


def test_xlsx_hidden_selected_sheet_blocks_import():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["id", "ref", "weight", "date"])
    sheet.append(["S1", "R1", 10, "30/09/2026"])
    workbook.create_sheet("Visible")
    sheet.sheet_state = "hidden"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(ImportErrorDetail, match="hoja seleccionada está oculta"):
        import_data(output.getvalue(), "hidden.xlsx", mapping(sheet="Data"))


@pytest.mark.parametrize("hostile_name", ["xl/worksheets/sheet1.xml", "../outside.txt"])
def test_xlsx_duplicate_or_traversal_entry_is_rejected(hostile_name):
    valid = xlsx_data([["id", "ref", "weight", "date"], ["S1", "R1", 10, "30/09/2026"]])
    altered = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(valid)) as source, zipfile.ZipFile(altered, "w") as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            target.writestr(hostile_name, b"ambiguous")
    with pytest.raises(ImportErrorDetail, match="entradas ZIP repetidas o rutas"):
        import_data(altered.getvalue(), "hostile.xlsx", mapping())


def test_xlsx_external_entity_never_becomes_imported_value(tmp_path):
    witness = tmp_path / "secret.txt"
    witness.write_text("SECRET_WITNESS", encoding="utf-8")
    valid = xlsx_data([["id", "ref", "weight", "date"], ["S1", "R1", 10, "30/09/2026"]])
    altered = io.BytesIO()
    doctype = f'<!DOCTYPE worksheet [<!ENTITY xxe SYSTEM "{witness.as_uri()}">]>'.encode()
    with zipfile.ZipFile(io.BytesIO(valid)) as source, zipfile.ZipFile(altered, "w") as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<worksheet ", doctype + b"<worksheet ", 1)
                content = content.replace(b"<v>10</v>", b"<v>&xxe;</v>", 1)
                assert b"&xxe;" in content
            target.writestr(info, content)
    with pytest.raises(ImportErrorDetail) as error:
        import_data(altered.getvalue(), "entity.xlsx", mapping())
    assert "SECRET_WITNESS" not in str(error.value)
    assert witness.read_text(encoding="utf-8") == "SECRET_WITNESS"


def test_csv_argentine_numbers_zeroes_and_provenance():
    data = csv_data("S1;0000123;1.234,56;30/09/2026\n")
    result = import_data(data, "origen.csv", mapping())
    assert result["accepted"] == 1 and not result["rejected"]
    record = result["records"][0]
    assert record["reference"] == "0000123"
    assert record["attributes"]["weight"]["value"] == "1234.56"
    source = record["provenance"]["attributes.weight"]
    assert source["document"] == bytes_hash(data)
    assert (source["row"], source["column"], source["raw"]) == (2, "weight", "1.234,56")


@pytest.mark.parametrize("extension", ["xlsx", "xls"])
@pytest.mark.parametrize("mac_epoch", [False, True])
def test_excel_fictitious_leap_day_never_aliases_real_date(extension, mac_epoch):
    # 59 and 60 alias to February 28 in common readers for the 1900 system.
    # Only the nonexistent day must be rejected; the 1904 epoch has no such defect.
    rows = [["id", "ref", "weight", "date"], *[[f"S{n}", f"R{n}", 1, n] for n in (59, 60, 61)]]
    output = io.BytesIO()
    if extension == "xlsx":
        from openpyxl.utils.datetime import CALENDAR_MAC_1904

        workbook = Workbook()
        if mac_epoch:
            workbook.epoch = CALENDAR_MAC_1904
        sheet = workbook.active
        for row in rows:
            sheet.append(row)
        for index in range(2, 5):
            sheet.cell(index, 4).number_format = "yyyy-mm-dd"
        workbook.save(output)
        workbook.close()
    else:
        workbook = xlwt.Workbook()
        workbook.dates_1904 = mac_epoch
        sheet = workbook.add_sheet("Data")
        style = xlwt.easyxf(num_format_str="YYYY-MM-DD")
        for row_index, row in enumerate(rows):
            for col_index, value in enumerate(row):
                sheet.write(
                    row_index,
                    col_index,
                    value,
                    style if col_index == 3 and row_index else xlwt.Style.default_style,
                )
        workbook.save(output)
    result = import_data(output.getvalue(), "dates." + extension, mapping(allow_xls_cached_values=True))
    dates = {r["id"]: r["attributes"]["service_date"]["value"] for r in result["records"]}
    if mac_epoch:
        assert dates == {"S59": "1904-02-29", "S60": "1904-03-01", "S61": "1904-03-02"}
        assert not result["rejected"]
    else:
        assert dates == {"S59": "1900-02-28", "S61": "1900-03-01"}
        assert len(result["rejected"]) == 1
        assert result["rejected"] == [3]
        assert "29/02/1900" in str(result["issues"])


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("1.234,56", "1234.56"),
        ("-1.000,01", "-1000.01"),
        ("0001,20", "1.2"),
        ("0,10", "0.1"),
        ("1.234", "1234"),
        ("1", "1"),
        ("-0,00", "0"),
    ],
)
def test_argentine_decimal(raw, expected):
    assert parse_decimal(raw, ",", ".") == expected


@pytest.mark.parametrize(
    "raw",
    ["1.23,45", "12.34", "1,000.25", "1 000,20", "1,2,3", "NaN", "+1", "1e2", "(100)", "1.234.56", "--1", ""],
)
def test_unsafe_numeric_strings_rejected(raw):
    with pytest.raises(ValueError):
        parse_decimal(raw, ",", ".")


def test_explicit_english_format():
    assert parse_decimal("1,234.56", ".", ",") == "1234.56"
    assert parse_decimal("1.234", ".", ",") == "1.234"


def test_invalid_date_has_row_field_and_original():
    result = import_data(csv_data("S1;0001;20;31/13/2026\n"), "bad.csv", mapping())
    assert result["accepted"] == 0 and result["rejected"] == [2]
    issue = result["issues"][0]
    assert issue["field"] == "attributes.service_date" and issue["raw"] == "31/13/2026"
    assert "Fila 2" in issue["message"]


def test_ambiguous_date_is_rejected():
    config = mapping(date_formats=["%d/%m/%Y", "%m/%d/%Y"])
    result = import_data(csv_data("S1;0001;20;01/02/2026\n"), "bad.csv", config)
    assert result["rejected"] == [2]
    assert "ambigua" in result["issues"][0]["message"]


def test_locale_month_names_not_accepted():
    with pytest.raises(ValidationError):
        mapping(date_formats=["%d %B %Y"])


def test_xlsx_formula_never_evaluated_or_taken_as_cache():
    data = xlsx_data([["id", "ref", "weight", "date"], ["S1", "00001", "=10*2", "01/09/2026"]])
    result = import_data(data, "data.xlsx", mapping())
    assert result["accepted"] == 0
    assert "fórmula" in result["issues"][0]["message"]


def test_xlsx_numeric_id_requires_explicit_policy():
    data = xlsx_data(
        [["id", "ref", "weight", "date"], ["S1", 123, 20, datetime(2026, 9, 1)]], {"B2": "000000"}
    )
    assert import_data(data, "data.xlsx", mapping())["accepted"] == 0
    config = mapping().model_dump()
    config["columns"][1]["numeric_text"] = "formatted"
    result = import_data(data, "data.xlsx", ImportMapping.model_validate(config))
    assert result["records"][0]["reference"] == "000123"


def test_timestamps_are_not_silently_truncated():
    data = xlsx_data([["id", "ref", "weight", "date"], ["S1", "0001", 20, datetime(2026, 9, 1, 12, 30)]])
    result = import_data(data, "data.xlsx", mapping())
    assert result["rejected"] == [2]
    assert "hora" in result["issues"][0]["message"]


def test_xlsx_numeric_cell_ignores_text_locale():
    data = xlsx_data([["id", "ref", "weight", "date"], ["S1", "0001", 1234.56, datetime(2026, 9, 1)]])
    result = import_data(data, "data.xlsx", mapping())
    assert result["records"][0]["attributes"]["weight"]["value"] == "1234.56"


def test_multisheet_requires_selection():
    wb = Workbook()
    wb.create_sheet("Segundo")
    data = io.BytesIO()
    wb.save(data)
    with pytest.raises(ImportErrorDetail, match="varias hojas"):
        import_data(data.getvalue(), "sheets.xlsx", mapping())


@pytest.mark.parametrize("headers", ["id;ref;weight;date;weight", "id;ref;date"])
def test_ambiguous_or_missing_headers_fail_file(headers):
    with pytest.raises(ImportErrorDetail, match="columna"):
        import_data((headers + "\nS1;0001;10;01/01/2026\n").encode(), "bad.csv", mapping())


def test_reordered_columns_same_normalized_records():
    first = import_data(csv_data("S1;0001;10;01/01/2026\n"), "one.csv", mapping())["records"][0]
    second = import_data(b"date;weight;ref;id\n01/01/2026;10;0001;S1\n", "two.csv", mapping())["records"][0]
    first.pop("provenance")
    second.pop("provenance")
    assert first == second


def test_duplicate_business_ids_visible_not_silently_deduplicated():
    result = import_data(csv_data("S1;0001;10;01/01/2026\nS1;0002;20;01/01/2026\n"), "repeat.csv", mapping())
    assert result["accepted"] == 1 and result["rejected"] == [3]
    assert "repetido" in result["issues"][0]["message"]


def test_unknown_concept_not_canonicalized_by_accident():
    config = ImportMapping.model_validate(
        {
            "id": "c",
            "version": "1",
            "entity": "charges",
            "concept_map": {"Freight": "base"},
            "columns": [
                constant("id", "C1"),
                constant("reference", "S1"),
                constant("carrier", "C"),
                constant("agreement", "A"),
                constant("settlement", "L"),
                constant("currency", "ARS"),
                column("concept", "concept"),
                column("amount", "amount", "decimal"),
            ],
        }
    )
    result = import_data(b"concept;amount\nbase;100\n", "charges.csv", config)
    assert result["records"][0]["concept"] == "external:base"
    assert result["issues"][0]["category"] == "warning"


def test_xls_requires_explicit_cached_value_acceptance():
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Data")
    for r, row in enumerate([["id", "ref", "weight", "date"], ["S1", "0001", 12.5, "01/09/2026"]]):
        for c, value in enumerate(row):
            sheet.write(r, c, value)
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(ImportErrorDetail, match="valores guardados"):
        import_data(output.getvalue(), "legacy.xls", mapping())
    result = import_data(output.getvalue(), "legacy.xls", mapping(allow_xls_cached_values=True))
    assert result["accepted"] == 1 and result["issues"][0]["category"] == "warning"


@given(st.binary(min_size=1, max_size=500))
@settings(max_examples=40)
def test_malformed_xlsx_fuzz_is_bounded_and_classified(data):
    with pytest.raises(ImportErrorDetail):
        import_data(data, "malformed.xlsx", mapping())


@given(st.text(max_size=40))
@settings(max_examples=100)
def test_numeric_fuzz_never_crashes_or_returns_nonfinite(text):
    try:
        result = parse_decimal(text, ",", ".")
    except ValueError:
        return
    from decimal import Decimal

    assert Decimal(result).is_finite()


def test_boolean_ambiguity_rejected():
    config = ColumnMapping(
        source="flag", target="attributes.flag", type="boolean", true_values=["yes"], false_values=["yes"]
    )
    with pytest.raises(ValueError, match="booleano único"):
        convert(Cell("yes"), config, mapping())


def workbook_with_numeric_literal(literal):
    import zipfile

    data = xlsx_data([["id", "ref", "weight", "date"], ["S1", "0001", 0.1, "01/09/2026"]])
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as origin, zipfile.ZipFile(output, "w") as target:
        for name in origin.namelist():
            content = origin.read(name)
            if name == "xl/worksheets/sheet1.xml":
                assert b"<v>0.1</v>" in content
                content = content.replace(b"<v>0.1</v>", ("<v>" + literal + "</v>").encode())
            target.writestr(name, content)
    return output.getvalue()


def test_xlsx_preserves_original_decimal_literal_without_binary_float_loss():
    literal = "1000000000000000.01"
    result = import_data(workbook_with_numeric_literal(literal), "exact.xlsx", mapping())
    assert result["accepted"] == 1
    record = result["records"][0]
    assert record["attributes"]["weight"]["value"] == literal
    assert record["provenance"]["attributes.weight"]["raw"] == literal


def test_xlsx_unbounded_exponent_is_rejected_before_decimal_expansion():
    result = import_data(workbook_with_numeric_literal("1e999999"), "extreme.xlsx", mapping())
    assert result["accepted"] == 0 and result["rejected"] == [2]


@pytest.mark.parametrize(
    "literal, expected", [("1.2345e+2", "123.45"), ("1.2345e-2", "0.012345"), ("1e-12", "0.000000000001")]
)
def test_xlsx_scientific_notation_remains_exact(literal, expected):
    result = import_data(workbook_with_numeric_literal(literal), "scientific.xlsx", mapping())
    assert result["accepted"] == 1
    assert result["records"][0]["attributes"]["weight"]["value"] == expected
    assert result["records"][0]["provenance"]["attributes.weight"]["raw"] == literal


def test_original_fractional_numeric_identifier_cannot_become_integer_after_float_rounding():
    config = ColumnMapping(source="ref", target="reference", type="text", numeric_text="formatted")
    cell = Cell(99999999999999.0, numeric_literal="99999999999999.001")
    with pytest.raises(ValueError, match="contiene decimales"):
        convert(cell, config, mapping())


def test_original_numeric_tokens_follow_sparse_sheet_coordinates():
    data = xlsx_data([["id", "ref", "weight", "date"], [], ["S1", "0001", 12.5, "01/09/2026"]])
    result = import_data(data, "sparse.xlsx", mapping())
    assert result["accepted"] == 1
    provenance = result["records"][0]["provenance"]["attributes.weight"]
    assert provenance["row"] == 3 and provenance["raw"] == "12.5"


def test_xlsx_populated_cell_under_blank_header_is_rejected():
    data = xlsx_data([["id", "ref", "weight", "date"], ["S1", "0001", 12.5, "01/09/2026", "EXTRA"]])
    result = import_data(data, "extra.xlsx", mapping())
    assert result["accepted"] == 0
    assert result["rejected"] == [2]
    assert any("fuera de las columnas declaradas" in issue["message"] for issue in result["issues"])
