"""Limits reject entire work, and file-backed outputs retain exact evidence."""

import asyncio
import json
import sqlite3
import zipfile
from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from freight_audit import importing, reporting
from freight_audit.canonical import canonical
from freight_audit.engine import audit
from freight_audit.models import AuditResult
from freight_audit.storage import IntegrityError, Store


def saved(tmp_path, make_dataset):
    store = Store(tmp_path / "resources.db")
    dataset = make_dataset()
    result = audit(dataset)
    return store, store.save(dataset, result)


def test_file_workbook_keeps_exact_cells_styles_and_filters(tmp_path, make_dataset):
    store, identifier = saved(tmp_path, make_dataset)
    run = store.load(identifier)
    target = tmp_path / "streamed.xlsx"
    with target.open("wb") as output:
        reporting.write_workbook(run, output)
    book = load_workbook(target)
    assert book["Hallazgos"]["F2"].value == "100"
    assert book["Hallazgos"]["F2"].data_type == "s"
    assert book["Hallazgos"].freeze_panes == "A2"
    assert book["Hallazgos"].auto_filter.ref == "A1:P2"
    assert book["Hallazgos"]["F2"].alignment.wrap_text
    assert book["Hallazgos"]["A1"].font.bold
    assert len(book.sheetnames) == 8
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    book.close()


def test_impossible_bundle_budget_rejects_before_workbook(tmp_path, make_dataset, monkeypatch):
    store, identifier = saved(tmp_path, make_dataset)
    monkeypatch.setattr(reporting, "MAX_BUNDLE_UNCOMPRESSED_BYTES", 200)

    def unexpected(*_args, **_kwargs):
        raise AssertionError("A package already over budget must not construct Excel")

    monkeypatch.setattr(reporting, "workbook_bytes", unexpected)
    monkeypatch.setattr(reporting, "write_workbook", unexpected, raising=False)
    with pytest.raises(reporting.ReportLimitError, match="límite"):
        reporting.bundle_bytes(store, identifier)


def test_file_bundle_verifies_sources_and_complete_canonical_output(tmp_path, make_dataset):
    store, identifier = saved(tmp_path, make_dataset)
    target = tmp_path / "streamed.zip"
    with target.open("w+b") as output:
        reporting.write_bundle(store, identifier, output)
    verified = reporting.verify_bundle(target)
    assert verified["id"] == identifier
    with zipfile.ZipFile(target) as archive:
        assert json.loads(archive.read("audit.json"))["result_hash"] == store.load(identifier)["result_hash"]
        assert {"auditoria.xlsx", "manifest.json", "reporte.html"} <= set(archive.namelist())


def test_history_list_rejects_altered_amounts_before_display(tmp_path, make_dataset):
    store, identifier = saved(tmp_path, make_dataset)
    result = deepcopy(store.load(identifier)["result"])
    result["summary"]["currencies"]["ARS"]["actual"] = "999"
    with sqlite3.connect(store.path) as connection:
        connection.execute("DROP TRIGGER immutable_runs_UPDATE")
        connection.execute("UPDATE runs SET result=? WHERE id=?", (canonical(result), identifier))
    with pytest.raises(IntegrityError):
        store.list_runs()


def test_persistence_independently_checks_input_totals(tmp_path, make_dataset, monkeypatch):
    dataset = make_dataset()
    tampered = audit(dataset).model_dump(mode="json")
    tampered["summary"]["currencies"]["ARS"]["actual"] = "999"
    result = AuditResult.model_validate(tampered)
    # Simulate a systematic summarizer bug that repeating the motor would share.
    monkeypatch.setattr("freight_audit.storage.audit", lambda _: result)
    store = Store(tmp_path / "totals.db")
    with pytest.raises(IntegrityError):
        store.save(dataset, result)
    assert store.list_runs() == []


def test_normalization_budget_aborts_without_partial_records(monkeypatch):
    from test_importing import csv_data, mapping

    monkeypatch.setattr(importing, "MAX_NORMALIZED_BYTES", 2000, raising=False)
    data = csv_data("".join(f"S{i};{i:05};10;01/05/2026\n" for i in range(20)))
    with pytest.raises(importing.ImportErrorDetail, match="normalizad"):
        importing.import_data(data, "bounded.csv", mapping())


def test_cell_budget_stops_expansion_even_with_small_csv(monkeypatch):
    from test_importing import csv_data, mapping

    monkeypatch.setattr(importing, "MAX_TABULAR_CELLS", 10, raising=False)
    data = csv_data("".join(f"S{i};{i:05};10;01/05/2026\n" for i in range(4)))
    with pytest.raises(importing.ImportErrorDetail, match="celdas"):
        importing.import_data(data, "small-but-wide.csv", mapping())


def test_incremental_json_has_identical_exact_utf8_bytes(make_dataset):
    from freight_audit.canonical import canonical_chunks

    value = {"text": "áñ < >", "number": Decimal("100.2500"), "nested": [False, None, make_dataset()]}
    assert b"".join(canonical_chunks(value)) == canonical(value).encode()
    assert b'"100.25"' in b"".join(canonical_chunks(value))


def test_temporary_download_is_cleaned_when_client_disconnects(tmp_path):
    from starlette.requests import ClientDisconnect

    from freight_audit.server import TemporaryArtifactResponse

    path = tmp_path / "temporary-report.txt"
    path.write_bytes(b"private temporary output")
    response = TemporaryArtifactResponse(path, cleanup=path.unlink)

    async def send(message):
        if message["type"] == "http.response.body":
            raise OSError("Simulated disconnected client")

    async def receive():
        return {"type": "http.disconnect"}

    with pytest.raises((OSError, ClientDisconnect)):
        asyncio.run(
            response({"type": "http", "method": "GET", "headers": [], "extensions": {}}, receive, send)
        )
    assert not path.exists()


def test_matching_budget_rejects_before_expanding_candidates(make_dataset, monkeypatch):
    from freight_audit import engine

    monkeypatch.setattr(engine, "MAX_AUDIT_WORK", 1, raising=False)

    def unexpected(*_args, **_kwargs):
        raise AssertionError("Over-budget input must not construct candidate traces")

    monkeypatch.setattr(engine, "inspect_matching", unexpected)
    with pytest.raises(ValueError, match="presupuesto"):
        engine.audit(make_dataset())
    with pytest.raises(ValueError, match="presupuesto"):
        engine.matching_preflight(make_dataset())


def test_temporary_download_streams_even_when_pathsend_is_advertised(tmp_path):
    from freight_audit.server import TemporaryArtifactResponse

    path = tmp_path / "temporary.txt"
    path.write_bytes(b"complete artifact")
    response = TemporaryArtifactResponse(path, cleanup=path.unlink)
    bodies = []

    async def send(message):
        assert message["type"] != "http.response.pathsend", (
            "Cleanup must wait for transfer, not a deferred pathname"
        )
        if message["type"] == "http.response.body":
            bodies.append(message["body"])

    async def receive():
        return {"type": "http.disconnect"}

    asyncio.run(
        response(
            {"type": "http", "method": "GET", "headers": [], "extensions": {"http.response.pathsend": {}}},
            receive,
            send,
        )
    )
    assert b"".join(bodies) == b"complete artifact"
    assert not path.exists()


def test_xlsx_budget_returns_warning_in_bundle_without_truncating_json(tmp_path, make_dataset, monkeypatch):
    store, identifier = saved(tmp_path, make_dataset)
    monkeypatch.setattr(reporting, "MAX_XLSX_XML_BUDGET", 1, raising=False)
    with pytest.raises(reporting.ReportLimitError, match="presupuesto"):
        reporting.workbook_bytes(store.load(identifier))
    bundle = reporting.bundle_bytes(store, identifier)
    with zipfile.ZipFile(__import__("io").BytesIO(bundle)) as archive:
        assert "auditoria.xlsx" not in archive.namelist()
        assert "ADVERTENCIA_EXPORTACION.txt" in archive.namelist()
    assert reporting.verify_bundle(bundle)["result"] == store.load(identifier)["result"]
