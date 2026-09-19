import io
import sqlite3
import zipfile
from concurrent.futures import ThreadPoolExecutor

import pytest
from openpyxl import load_workbook

from freight_audit.engine import audit
from freight_audit.models import Dataset, Decision
from freight_audit.reporting import bundle_bytes, html_report, verify_bundle, workbook_bytes
from freight_audit.storage import IntegrityError, Store


def saved(tmp_path, make_dataset):
    store = Store(tmp_path / "data.db")
    dataset = make_dataset()
    result = audit(dataset)
    run_id = store.save(dataset, result)
    return store, run_id, result


def test_save_replay_and_idempotency(tmp_path, make_dataset):
    store, run_id, result = saved(tmp_path, make_dataset)
    assert store.save(make_dataset(), result) == run_id
    assert len(store.list_runs()) == 1
    assert store.replay(run_id)["identical"]


def test_human_decision_never_rewrites_finding(tmp_path, make_dataset):
    store, run_id, result = saved(tmp_path, make_dataset)
    before = store.load(run_id)
    store.add_decision(
        run_id,
        Decision(
            finding_id=result.findings[0].id,
            actor="Operador de prueba",
            action="REJECTED",
            note="Una decisión diferente al motor",
        ),
    )
    after = store.load(run_id)
    assert after["result"] == before["result"]
    assert after["result_hash"] == before["result_hash"]
    assert len(after["decisions"]) == 1
    assert store.replay(run_id)["identical"]


def test_decision_chain_serializes_concurrent_writers(tmp_path, make_dataset):
    store, run_id, result = saved(tmp_path, make_dataset)
    decision = Decision(
        finding_id=result.findings[0].id, actor="Prueba", action="APPROVED", note="Resolución de prueba"
    )
    # Lightweight database concurrency only, not parallel builds or workers.
    with ThreadPoolExecutor(max_workers=2) as executor:
        list(executor.map(lambda _: store.add_decision(run_id, decision), range(6)))
    assert len(store.decisions(run_id)) == 6


def test_changed_agreement_preserves_history(tmp_path, make_dataset):
    store, run_id, result = saved(tmp_path, make_dataset)
    raw = make_dataset().model_dump()
    raw["agreements"][0]["versions"][0]["rules"][0]["expression"]["value"]["value"] = "200"
    changed = Dataset.model_validate(raw)
    other = store.save(changed, audit(changed))
    assert other != run_id
    assert store.load(run_id)["result"]["findings"][0]["status"] == "PASS"
    assert store.load(other)["result"]["findings"][0]["status"] == "FAIL"


def test_db_update_delete_blocked(tmp_path, make_dataset):
    store, run_id, _ = saved(tmp_path, make_dataset)
    with store.connect() as conn:
        with pytest.raises(sqlite3.IntegrityError):
            conn.execute("UPDATE runs SET result=? WHERE id=?", ("{}", run_id))
        with pytest.raises(sqlite3.IntegrityError):
            conn.execute("DELETE FROM runs WHERE id=?", (run_id,))


def test_tampered_result_detected_after_trigger_bypass(tmp_path, make_dataset):
    store, run_id, _ = saved(tmp_path, make_dataset)
    with store.connect() as conn:
        conn.execute("DROP TRIGGER immutable_runs_UPDATE")
        conn.execute("UPDATE runs SET result=? WHERE id=?", ("{}", run_id))
    with pytest.raises(IntegrityError):
        store.load(run_id)


def test_tampered_source_detected(tmp_path):
    store = Store(tmp_path / "data.db")
    source = store.put_source(b"original")
    with store.connect() as conn:
        conn.execute("DROP TRIGGER immutable_sources_UPDATE")
        conn.execute("UPDATE sources SET data=? WHERE hash=?", (b"changed", source))
    with pytest.raises(IntegrityError):
        store.source(source)


def test_tampered_decision_detected(tmp_path, make_dataset):
    store, run_id, result = saved(tmp_path, make_dataset)
    store.add_decision(
        run_id,
        Decision(
            finding_id=result.findings[0].id, actor="Prueba", action="APPROVED", note="Resolución de prueba"
        ),
    )
    with store.connect() as conn:
        conn.execute("DROP TRIGGER immutable_decisions_UPDATE")
        conn.execute("UPDATE decisions SET payload=?", ("{}",))
    with pytest.raises(IntegrityError):
        store.decisions(run_id)


def test_false_result_cannot_be_saved(tmp_path, make_dataset):
    store = Store(tmp_path / "data.db")
    dataset = make_dataset()
    raw = audit(dataset).model_dump()
    raw["findings"][0]["status"] = "FAIL"
    from freight_audit.models import AuditResult

    with pytest.raises(IntegrityError):
        store.save(dataset, AuditResult.model_validate(raw))


def test_engine_artifact_change_blocks_replay(tmp_path, make_dataset, monkeypatch):
    store, run_id, _ = saved(tmp_path, make_dataset)
    monkeypatch.setattr("freight_audit.storage.engine_artifact_hash", lambda: "f" * 64)
    with pytest.raises(IntegrityError, match="código"):
        store.replay(run_id)


def test_newer_database_schema_is_not_opened(tmp_path):
    path = tmp_path / "future.db"
    with sqlite3.connect(path) as conn:
        conn.execute("PRAGMA user_version=99")
    with pytest.raises(IntegrityError):
        Store(path)


def test_migration_reopening_and_consistent_backup(tmp_path, make_dataset):
    store, run_id, _ = saved(tmp_path, make_dataset)
    assert Store(store.path).replay(run_id)["identical"]
    backup = tmp_path / "backup.db"
    store.backup(backup)
    assert Store(backup).replay(run_id)["identical"]
    with pytest.raises(ValueError):
        store.backup(backup)
    with store.connect() as conn:
        assert conn.execute("PRAGMA user_version").fetchone()[0] == 1


def test_workbook_is_real_and_has_exact_money_and_no_formulas(tmp_path, make_dataset):
    store, run_id, _ = saved(tmp_path, make_dataset)
    book = load_workbook(io.BytesIO(workbook_bytes(store.load(run_id))), data_only=False)
    assert book.sheetnames == [
        "Resumen",
        "Hallazgos",
        "Cálculos",
        "Problemas de datos",
        "Evidencia",
        "Decisiones",
        "Origen de datos",
        "Metadata",
    ]
    assert book["Hallazgos"]["F2"].value == "100"
    assert book["Hallazgos"]["F2"].data_type == "s"
    assert book["Hallazgos"].freeze_panes == "A2"
    for sheet in book:
        assert sheet.auto_filter.ref
        assert all(cell.data_type != "f" for row in sheet for cell in row)
    book.close()


def test_report_escapes_untrusted_html_and_excel_formula(tmp_path, raw_dataset):
    raw_dataset["label"] = "<script>alert(1)</script>"
    raw_dataset["charges"][0]["concept"] = '=HYPERLINK("https://evil", "click")'
    dataset = Dataset.model_validate(raw_dataset)
    store = Store(tmp_path / "data.db")
    run_id = store.save(dataset, audit(dataset))
    run = store.load(run_id)
    html = html_report(run)
    assert "<script>" not in html and "&lt;script&gt;" in html
    book = load_workbook(io.BytesIO(workbook_bytes(run)))
    assert book["Hallazgos"]["D2"].data_type == "s"
    book.close()


def test_portable_bundle_integrity(tmp_path, make_dataset):
    store, run_id, _ = saved(tmp_path, make_dataset)
    data = bundle_bytes(store, run_id)
    assert verify_bundle(data)["id"] == run_id
    original = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as modified:
        for name in original.namelist():
            modified.writestr(
                name, b"changed" if name in {"snapshot.json", "audit.json"} else original.read(name)
            )
    with pytest.raises(IntegrityError):
        verify_bundle(output.getvalue())


def test_no_unknown_decision_or_evidence(tmp_path, make_dataset):
    store, run_id, result = saved(tmp_path, make_dataset)
    with pytest.raises(ValueError):
        store.add_decision(
            run_id, Decision(finding_id="missing", actor="Prueba", action="APPROVED", note="No existe")
        )
    with pytest.raises(ValueError):
        store.add_decision(
            run_id,
            Decision(
                finding_id=result.findings[0].id,
                actor="Prueba",
                action="APPROVED",
                note="No existe evidencia",
                evidence_ids=["missing"],
            ),
        )


def test_source_update_requires_process_restart(tmp_path, make_dataset, monkeypatch):
    store = Store(tmp_path / "data.db")
    dataset = make_dataset()
    monkeypatch.setattr("freight_audit.storage.engine_artifact_hash", lambda: "e" * 64)
    with pytest.raises(IntegrityError, match="reiniciar"):
        store.save(dataset, audit(dataset))


def test_excel_long_cell_is_not_silently_truncated(tmp_path, make_dataset):
    from freight_audit.reporting import ReportLimitError

    store, run_id, _ = saved(tmp_path, make_dataset)
    run = store.load(run_id)
    run["result"]["findings"][0]["reasons"] = ["x" * 32768]
    with pytest.raises(ReportLimitError, match="sin truncar"):
        workbook_bytes(run)


def test_excel_row_limit_preserves_full_bundle_json(tmp_path, make_dataset, monkeypatch):
    store, run_id, _ = saved(tmp_path, make_dataset)
    # Make the real row-limit branch reachable with a normal complete audit.
    monkeypatch.setattr("freight_audit.reporting.MAX_XLSX_ROWS", 2)
    bundle = bundle_bytes(store, run_id)
    with zipfile.ZipFile(io.BytesIO(bundle)) as archive:
        assert "auditoria.xlsx" not in archive.namelist()
        assert "ADVERTENCIA_EXPORTACION.txt" in archive.namelist()
    assert verify_bundle(bundle)["id"] == run_id
