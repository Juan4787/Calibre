"""Explicit reusable CSV/XLSX/XLS mapping, with cell-level provenance and rejects."""

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from dataclasses import field as dataclass_field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Literal

import openpyxl
import xlrd
from defusedxml.ElementTree import iterparse
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from pydantic import Field, ValidationError, model_validator

from .canonical import bytes_hash, decimal_text, number
from .models import Charge, DataIssue, Identifier, Model, Shipment, SourceRef, Value

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 150 * 1024 * 1024
MAX_ROWS = 200001
MAX_COLUMNS = 250


class ImportErrorDetail(ValueError):
    pass


class ColumnMapping(Model):
    source: str | None = None
    target: Identifier
    type: Literal["text", "decimal", "date", "boolean"]
    unit: str | None = None
    required: bool = True
    constant: str | bool | None = None
    numeric_text: Literal["reject", "formatted"] = "reject"
    true_values: list[str] = Field(default_factory=lambda: ["true"])
    false_values: list[str] = Field(default_factory=lambda: ["false"])

    @model_validator(mode="after")
    def source_or_constant(self):
        if (self.source is None) == (self.constant is None):
            raise ValueError("Cada campo necesita una columna o un valor constante, exclusivamente.")
        return self


class ImportMapping(Model):
    schema_version: Literal[1] = 1
    id: Identifier
    version: Identifier
    entity: Literal["shipments", "charges"]
    sheet: str | None = None
    header_row: int = Field(default=1, ge=1, le=1000)
    delimiter: Literal[",", ";", "\t", "|"] = ";"
    encoding: Literal["utf-8-sig", "cp1252"] = "utf-8-sig"
    decimal_separator: Literal[".", ","] = ","
    thousands_separator: Literal[".", ",", "", " "] = "."
    date_formats: list[str] = Field(default_factory=lambda: ["%d/%m/%Y"])
    columns: list[ColumnMapping] = Field(min_length=1)
    concept_map: dict[str, str] = Field(default_factory=dict)
    allow_xls_cached_values: bool = False

    @model_validator(mode="after")
    def valid_mapping(self):
        if self.decimal_separator == self.thousands_separator:
            raise ValueError("Los separadores decimal y de miles deben ser diferentes.")
        targets = [c.target for c in self.columns]
        if len(set(targets)) != len(targets):
            raise ValueError("Dos columnas no pueden escribir el mismo campo.")
        required = (
            {"id", "reference", "carrier"}
            if self.entity == "shipments"
            else {"id", "reference", "carrier", "agreement", "settlement", "concept", "amount", "currency"}
        )
        if not required <= set(targets):
            raise ValueError(f"Faltan campos obligatorios: {', '.join(sorted(required - set(targets)))}.")
        for column in self.columns:
            if column.target not in required and not column.target.startswith("attributes."):
                raise ValueError(
                    f"Campo desconocido '{column.target}'; los atributos adicionales usan attributes.nombre."
                )
            if column.target in required:
                expected_type = "decimal" if column.target == "amount" else "text"
                if column.type != expected_type:
                    raise ValueError(f"El campo '{column.target}' requiere tipo {expected_type}.")
        for fmt in self.date_formats:
            if (
                "%Y" not in fmt
                or "%m" not in fmt
                or "%d" not in fmt
                or any(token in fmt for token in ("%y", "%b", "%B", "%x", "%c"))
            ):
                raise ValueError("Las fechas requieren día, mes numérico y año de cuatro dígitos.")
        return self


@dataclass
class Cell:
    value: object
    format: str = ""
    formula: bool = False
    numeric_literal: str | None = None

    def raw(self):
        if self.numeric_literal is not None:
            return self.numeric_literal
        if isinstance(self.value, (date, datetime)):
            return self.value.isoformat()
        return "" if self.value is None else str(self.value)


@dataclass
class Tabular:
    sheets: list[str]
    sheet: str
    rows: list[list[Cell]]
    warnings: list[str]
    row_numbers: list[int] = dataclass_field(default_factory=list)


def preserve_xlsx_numbers(data: bytes, worksheet_path: str, rows: list[list[Cell]]) -> None:
    """Overlay the original numeric tokens, bounded to one XML row at a time.

    openpyxl supplies dates, formats and formulas; decimal values must bypass
    its binary float conversion. _worksheet_path belongs to ReadOnlyWorksheet
    and is covered by actual workbook import tests against the pinned version.
    """
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    sheet_data = None
    row_index = 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive, archive.open(worksheet_path.lstrip("/")) as xml:
        for event, element in iterparse(xml, events=("start", "end")):
            if event == "start" and element.tag == namespace + "sheetData":
                sheet_data = element
            if event != "end" or element.tag != namespace + "row":
                continue
            row_index = int(element.get("r", row_index + 1))
            if not 1 <= row_index <= len(rows):
                raise ImportErrorDetail("La numeración de filas del libro es inconsistente.")
            column_index = 0
            for source in element.findall(namespace + "c"):
                column_index += 1
                if source.get("r"):
                    column, coordinate_row = coordinate_from_string(source.get("r"))
                    if coordinate_row != row_index:
                        raise ImportErrorDetail("La ubicación de una celda del libro es inconsistente.")
                    column_index = column_index_from_string(column)
                if not 1 <= column_index <= len(rows[row_index - 1]):
                    raise ImportErrorDetail("El rango de columnas del libro es inconsistente.")
                cell = rows[row_index - 1][column_index - 1]
                if (
                    source.get("t", "n") == "n"
                    and not cell.formula
                    and isinstance(cell.value, (int, float))
                    and not isinstance(cell.value, bool)
                ):
                    token = source.findtext(namespace + "v")
                    if token is None:
                        raise ImportErrorDetail("Falta el valor original de una celda numérica.")
                    cell.numeric_literal = token
            element.clear()
            if sheet_data is not None:
                sheet_data.clear()


def excel_number(cell: Cell) -> Decimal:
    raw = cell.numeric_literal if cell.numeric_literal is not None else str(cell.value)
    if len(raw) > 128 or not re.fullmatch(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?", raw):
        raise ValueError("La celda no contiene un número decimal válido.")
    value = Decimal(raw)
    parts = value.as_tuple()
    if (
        not value.is_finite()
        or len(parts.digits) > 36
        or not isinstance(parts.exponent, int)
        or not -12 <= parts.exponent <= 36
        or value.adjusted() >= 36
    ):
        raise ValueError("El número excede el límite de 36 dígitos o 12 decimales; revisar el archivo.")
    # Bound the exponent before expanding scientific notation into fixed point.
    return number(format(value, "f"))


def read_table(data: bytes, filename: str, mapping: ImportMapping) -> Tabular:
    if len(data) > MAX_FILE_BYTES:
        raise ImportErrorDetail("El archivo supera 25 MB; dividirlo en períodos más pequeños.")
    suffix = Path(filename).suffix.lower()
    warnings: list[str] = []
    try:
        if suffix == ".csv":
            reader = csv.reader(
                io.StringIO(data.decode(mapping.encoding)), delimiter=mapping.delimiter, strict=True
            )
            rows = []
            physical_lines = []
            previous_line = 0
            for index, csv_row in enumerate(reader):
                if index >= MAX_ROWS or len(csv_row) > MAX_COLUMNS:
                    raise ImportErrorDetail("El archivo supera el límite de filas o columnas.")
                physical_lines.append(previous_line + 1)
                previous_line = reader.line_num
                rows.append([Cell(value) for value in csv_row])
            return Tabular(["CSV"], "CSV", rows, warnings, physical_lines)
        if suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(info.file_size for info in archive.infolist()) > MAX_EXPANDED_BYTES:
                    raise ImportErrorDetail("El libro expandido supera 150 MB; dividirlo antes de importar.")
            workbook = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            try:
                names = workbook.sheetnames
                if mapping.sheet is None and len(names) > 1:
                    raise ImportErrorDetail(
                        "El libro tiene varias hojas. Seleccionar una hoja explícitamente: "
                        + ", ".join(names)
                    )
                sheet_name = mapping.sheet or names[0]
                if sheet_name not in names:
                    raise ImportErrorDetail(
                        f"No existe la hoja '{sheet_name}'. Disponibles: {', '.join(names)}."
                    )
                sheet = workbook[sheet_name]
                if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                    raise ImportErrorDetail(
                        "La hoja supera el límite de filas o columnas; revisar su rango utilizado."
                    )
                rows = []
                for index, row in enumerate(sheet.iter_rows()):
                    if index >= MAX_ROWS or len(row) > MAX_COLUMNS:
                        raise ImportErrorDetail("La hoja supera el límite de importación.")
                    rows.append(
                        [
                            Cell(cell.value, cell.number_format or "", cell.data_type in {"f", "e"})
                            for cell in row
                        ]
                    )
                preserve_xlsx_numbers(data, sheet._worksheet_path, rows)
                return Tabular(names, sheet_name, rows, warnings)
            finally:
                workbook.close()
        if suffix == ".xls":
            if not mapping.allow_xls_cached_values:
                raise ImportErrorDetail(
                    "El formato XLS sólo ofrece valores guardados y puede contener fórmulas sin recalcular. Convertir a XLSX con valores verificados o habilitar expresamente allow_xls_cached_values en el mapping."
                )
            warnings.append(
                "XLS: se usan valores guardados; confirmar que las fórmulas estaban recalculadas antes de guardar."
            )
            workbook_xls = xlrd.open_workbook(file_contents=data, formatting_info=True, on_demand=True)
            try:
                names = workbook_xls.sheet_names()
                if mapping.sheet is None and len(names) > 1:
                    raise ImportErrorDetail("Seleccionar una hoja del libro XLS: " + ", ".join(names))
                name = mapping.sheet or names[0]
                if name not in names:
                    raise ImportErrorDetail(f"No existe la hoja '{name}'.")
                sheet_xls = workbook_xls.sheet_by_name(name)
                if sheet_xls.nrows > MAX_ROWS or sheet_xls.ncols > MAX_COLUMNS:
                    raise ImportErrorDetail("El libro XLS supera el límite de importación.")
                rows = []
                for row_index in range(sheet_xls.nrows):
                    xls_row: list[Cell] = []
                    for cell in sheet_xls.row(row_index):
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook_xls.datemode)
                        fmt = workbook_xls.format_map[
                            workbook_xls.xf_list[cell.xf_index].format_key
                        ].format_str
                        xls_row.append(Cell(value, fmt, cell.ctype == xlrd.XL_CELL_ERROR))
                    rows.append(xls_row)
                return Tabular(names, name, rows, warnings)
            finally:
                workbook_xls.release_resources()
        raise ImportErrorDetail("Formato no admitido. Usar CSV, XLSX o XLS.")
    except ImportErrorDetail:
        raise
    except Exception as exc:
        raise ImportErrorDetail(
            "No se pudo leer el archivo. Verificar formato, codificación y que el libro no esté dañado ni protegido con contraseña."
        ) from exc


def parse_decimal(raw: str, decimal_separator: str, thousands_separator: str) -> str:
    raw = raw.strip()
    sign = "-?"
    decimal = re.escape(decimal_separator)
    integer = r"\d+"
    if thousands_separator:
        grouping = re.escape(thousands_separator)
        integer = rf"(?:\d+|\d{{1,3}}(?:{grouping}\d{{3}})+)"
    if not re.fullmatch(rf"{sign}{integer}(?:{decimal}\d+)?", raw):
        raise ValueError(
            "Número inválido para los separadores configurados; revisar agrupadores y decimales."
        )
    if thousands_separator:
        raw = raw.replace(thousands_separator, "")
    raw = raw.replace(decimal_separator, ".")
    # Canonicalize leading zeros without a binary float conversion.
    from decimal import Decimal

    return decimal_text(
        number(
            format(Decimal(raw), "f").lstrip("+")
            if not re.match(r"^-?0\d", raw)
            else decimal_text(Decimal(raw))
        )
    )


def convert(cell: Cell, column: ColumnMapping, mapping: ImportMapping) -> str | bool:
    if cell.formula:
        raise ValueError("La celda contiene una fórmula o un error de Excel; aportar el valor verificado.")
    value = cell.value
    if column.type == "text":
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if column.numeric_text != "formatted":
                raise ValueError(
                    "El identificador está guardado como número y puede haber perdido ceros iniciales; convertir a texto o configurar el formato explícito."
                )
            numeric = excel_number(cell)
            if numeric != int(numeric) or len(str(abs(int(numeric)))) > 15:
                raise ValueError(
                    "El identificador numérico contiene decimales o excede la precisión de Excel."
                )
            text = str(int(numeric))
            if re.fullmatch(r"0+", cell.format):
                text = text.zfill(len(cell.format))
            return text
        if not isinstance(value, str):
            raise ValueError("El campo requiere texto.")
        return value.strip()
    if column.type == "decimal":
        if isinstance(value, bool) or isinstance(value, (date, datetime)):
            raise ValueError("El campo requiere un número decimal.")
        if isinstance(value, (int, float)):
            return decimal_text(excel_number(cell))
        return parse_decimal(str(value), mapping.decimal_separator, mapping.thousands_separator)
    if column.type == "date":
        if isinstance(value, datetime):
            if value.time().isoformat() != "00:00:00" or value.tzinfo:
                raise ValueError(
                    "La fecha incluye hora o zona horaria; definir primero la fecha contractual sin truncarla silenciosamente."
                )
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        matches = set()
        for fmt in mapping.date_formats:
            try:
                matches.add(datetime.strptime(str(value).strip(), fmt).date().isoformat())
            except ValueError:
                pass
        if len(matches) != 1:
            raise ValueError(
                "La fecha es inválida o ambigua para los formatos configurados; revisar día, mes y año."
            )
        return matches.pop()
    if isinstance(value, bool):
        return value
    true = str(value).strip() in column.true_values
    false = str(value).strip() in column.false_values
    if true == false:
        raise ValueError("El valor no identifica un booleano único según el mapping.")
    return true


def import_data(data: bytes, filename: str, mapping: ImportMapping) -> dict:
    table = read_table(data, filename, mapping)
    document = bytes_hash(data)
    if len(table.rows) < mapping.header_row:
        raise ImportErrorDetail("No existe la fila de encabezados seleccionada.")
    headers = [cell.raw().strip() for cell in table.rows[mapping.header_row - 1]]
    used = [column.source for column in mapping.columns if column.source]
    for source in used:
        if headers.count(source) != 1:
            raise ImportErrorDetail(
                f"La columna '{source}' aparece {headers.count(source)} veces; corregir encabezados o mapping."
            )
    positions = {header: index for index, header in enumerate(headers)}
    records = []
    issues = [DataIssue(category="warning", message=warning, document=document) for warning in table.warnings]
    rejected = []
    seen = set()
    preview: list[dict] = []
    for logical_row, row in enumerate(table.rows[mapping.header_row :], start=mapping.header_row + 1):
        row_number = table.row_numbers[logical_row - 1] if table.row_numbers else logical_row
        if all(cell.value is None or cell.value == "" for cell in row):
            continue
        if len(preview) < 10:
            preview.append(
                {
                    "row": row_number,
                    "values": {
                        header or f"Columna {i + 1}": row[i].raw() if i < len(row) else ""
                        for i, header in enumerate(headers)
                    },
                }
            )
        record: dict = {"attributes": {}, "provenance": {}}
        row_issues = []
        if len(row) > len(headers) and any(cell.value not in (None, "") for cell in row[len(headers) :]):
            row_issues.append(
                DataIssue(
                    category="row",
                    message=f"Fila {row_number}: hay valores fuera de las columnas declaradas; revisar separadores y encabezados.",
                    document=document,
                    row=row_number,
                )
            )
        for column in mapping.columns:
            index = positions.get(column.source, -1)
            cell = (
                Cell(column.constant)
                if column.source is None
                else row[index]
                if index < len(row)
                else Cell(None)
            )
            try:
                if cell.value is None or cell.value == "":
                    if column.required:
                        raise ValueError("Falta un valor obligatorio.")
                    continue
                value = convert(cell, column, mapping)
                if column.target == "concept" and mapping.concept_map:
                    if str(value) in mapping.concept_map:
                        value = mapping.concept_map[str(value)]
                    else:
                        issues.append(
                            DataIssue(
                                category="warning",
                                message=f"Fila {row_number}: concepto '{value}' sin equivalencia configurada; permanece sin clasificar.",
                                document=document,
                                row=row_number,
                                field=column.target,
                                raw=str(value),
                            )
                        )
                        value = "external:" + str(value)
                if column.target.startswith("attributes."):
                    record["attributes"][column.target[11:]] = Value(
                        type=column.type, value=value, unit=column.unit
                    )
                else:
                    record[column.target] = value
                record["provenance"][column.target] = SourceRef(
                    document=document,
                    filename=Path(filename).name,
                    sheet=table.sheet,
                    row=row_number,
                    column=column.source or "(constante del mapping)",
                    raw=cell.raw(),
                    transform=f"{mapping.id}@{mapping.version}:{column.type}",
                )
            except ValueError as exc:
                row_issues.append(
                    DataIssue(
                        category="row",
                        message=f"Fila {row_number}, campo '{column.target}', valor '{cell.raw()[:100]}': {exc}",
                        document=document,
                        row=row_number,
                        field=column.target,
                        raw=cell.raw()[:10000],
                    )
                )
        if not row_issues:
            try:
                model = (
                    Shipment.model_validate(record)
                    if mapping.entity == "shipments"
                    else Charge.model_validate(record)
                )
                if model.id in seen:
                    raise ValueError(f"Identificador '{model.id}' repetido; corregirlo en el archivo.")
                seen.add(model.id)
                records.append(model.model_dump(mode="json"))
            except (ValidationError, ValueError) as exc:
                message = (
                    "La fila no cumple el contrato de datos; revisar identificadores y campos obligatorios."
                    if isinstance(exc, ValidationError)
                    else str(exc)
                )
                row_issues.append(
                    DataIssue(
                        category="row",
                        message=f"Fila {row_number}: {message}",
                        document=document,
                        row=row_number,
                    )
                )
        if row_issues:
            rejected.append(row_number)
            issues.extend(row_issues)
    return {
        "document": document,
        "filename": Path(filename).name,
        "sheet": table.sheet,
        "sheets": table.sheets,
        "headers": headers,
        "preview": preview,
        "accepted": len(records),
        "rejected": rejected,
        "records": records,
        "issues": [issue.model_dump(mode="json") for issue in issues],
        "mapping": mapping.model_dump(mode="json"),
    }
