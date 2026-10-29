"""Exact-value workbooks and self-contained printable HTML, derived from traces."""

import hashlib
import html
import io
import re
import shutil
import stat
import tempfile
import zipfile
from itertools import chain
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter

from .canonical import canonical, canonical_chunks, digest, load_json
from .local_paths import checked_output_path
from .models import Dataset
from .storage import IntegrityError, Store, run_identity

STATUS_LABELS = {
    "PASS": "Coincide",
    "FAIL": "Discrepancia",
    "REVIEW": "Revisión humana",
    "UNDETERMINABLE": "Indeterminado",
}


MAX_XLSX_ROWS = 1_048_576
MAX_XLSX_CELL_CHARS = 32_767
MAX_XLSX_XML_BUDGET = 256 * 1024 * 1024
MAX_BUNDLE_UNCOMPRESSED_BYTES = 500 * 1024 * 1024


def provenance_note(run: dict) -> str:
    proof = run.get("metadata", {}).get("provenance")
    if not proof:
        return "Auditoría histórica sin registro de cotejo contra originales. La completitud documental requiere control independiente."
    note = (
        f"{proof['checked_records']} registros cotejados contra archivos y formatos conservados; "
        f"{proof['declared_records']} aportados sin cotejo. "
        "El cotejo no certifica autenticidad, contrato ni documentación completa del período."
    )
    control = run.get("metadata", {}).get("document_control")
    if control:
        note += " " + (
            "Control documental: inventario, cantidades e importes coinciden con el control aportado."
            if control["status"] == "matched"
            else "Control documental independiente pendiente."
        )
    return note


class ReportLimitError(ValueError):
    """Excel must not silently truncate authoritative data."""


def describe_trace(trace: dict) -> str:
    """Presentation comes exclusively from executed operands, never recomputed pricing."""
    operation = trace["op"]
    details = trace.get("details", {})
    output = trace.get("output")
    operands = [str(child.get("output", "?")) for child in trace.get("children", [])]
    symbols = {"add": " + ", "sub": " - ", "mul": " × ", "div": " ÷ "}
    if operation in symbols:
        return symbols[operation].join(operands) + " = " + str(output)
    if operation == "attr":
        return f"{details.get('field')}: {output}"
    if operation == "sum":
        return f"Suma de {details.get('field')} en {len(details.get('sources', []))} operaciones = {output}"
    if operation in {"lookup", "band"}:
        return f"Tabla {details.get('table')}, selección {canonical(details.get('selected', {}))}: {output}"
    if operation == "comparison":
        return f"Facturado {details.get('actual')} - esperado {details.get('expected')} = {output}; tolerancia {details.get('tolerance')}"
    if operation in {"round", "currency_round"}:
        return f"Redondeo a {details.get('scale')} decimales, {details.get('rounding')}: {output}"
    if operation == "version":
        return f"Versión {output}; fecha relevante: {details.get('date_field')}"
    if operation == "evidence":
        return f"Respaldos encontrados: {', '.join(details.get('found', [])) or 'ninguno'}; faltantes: {', '.join(details.get('missing', [])) or 'ninguno'}"
    if details.get("reason"):
        return str(details["reason"])
    return f"{operation}: {output if output is not None else canonical(details)}"


def trace_rows(trace, path=""):
    path = f"{path}/{trace['op']}"
    yield [path, describe_trace(trace), trace.get("output"), canonical(trace.get("details", {}))]
    for child in trace.get("children", []):
        yield from trace_rows(child, path)


def write_workbook(run: dict, output) -> None:
    """Write rows once to an owned stream; never retain a tree of Excel cells."""
    workbook = Workbook(write_only=True)
    alignment = Alignment(vertical="top", wrap_text=True)
    heading_fill = PatternFill("solid", fgColor="183E39")
    heading_font = Font(color="FFFFFF", bold=True)
    xml_budget = 0

    def cells(ws, values, heading=False):
        for value in values:
            cell = WriteOnlyCell(ws, value=value)
            cell.data_type = "s"
            cell.alignment = alignment
            if heading:
                cell.fill = heading_fill
                cell.font = heading_font
            yield cell

    def sheet(name, headings, rows):
        nonlocal xml_budget
        xml_budget += 64 * 1024
        ws = workbook.create_sheet(name)
        ws.freeze_panes = "A2"
        for col, title in enumerate(headings, 1):
            ws.column_dimensions[get_column_letter(col)].width = min(58, max(18, len(title) + 5))
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.append(cells(ws, headings, heading=True))
        row_count = 1
        for row in rows:
            values = [
                "" if value is None else ("Sí" if value else "No") if isinstance(value, bool) else str(value)
                for value in row
            ]
            if row_count >= MAX_XLSX_ROWS or any(len(value) > MAX_XLSX_CELL_CHARS for value in values):
                raise ReportLimitError(
                    "El detalle supera los límites de filas o de texto por celda de Excel. Exportar el paquete completo: audit.json conserva todos los datos sin truncar."
                )
            if any(
                ILLEGAL_CHARACTERS_RE.search(value)
                or any(ord(c) in {0xFFFE, 0xFFFF} or 0xD800 <= ord(c) <= 0xDFFF for c in value)
                for value in values
            ):
                raise ReportLimitError(
                    "Hay texto que Excel no puede representar. Exportar el paquete completo: "
                    "el JSON y los originales conservan los valores sin alterarlos."
                )
            # Conservative UTF-8/XML expansion plus cell/row markup; bound the
            # temporary worksheet files as well as the final compressed archive.
            xml_budget += sum(8 * len(value.encode("utf-8")) + 256 for value in values) + 128
            if xml_budget > MAX_XLSX_XML_BUDGET:
                raise ReportLimitError(
                    "La planilla supera el presupuesto de representación. Exportar el paquete completo: "
                    "el JSON y los originales conservan todos los datos sin truncar."
                )
            ws.append(cells(ws, values))
            row_count += 1
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headings))}{row_count}"
        return ws

    try:
        _workbook_sheets(run, sheet)
        # Own the ZIP context so an I/O failure also closes its archive handle.
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
            ExcelWriter(workbook, archive).write_data()
    finally:
        for ws in workbook.worksheets:
            if not ws.closed:
                ws.close()
            # openpyxl's pinned write-only writer uses temporary XML files.
            # Successful saves clean them; failed limits must do the same.
            writer = ws._writer
            if writer is not None and Path(writer.out).exists():
                writer.cleanup()
        workbook.close()


def _workbook_sheets(run: dict, sheet) -> None:
    summary = run["result"]["summary"]

    sheet(
        "Resumen",
        ["Moneda", "Métrica", "Valor exacto"],
        chain(
            (
                [currency, key, value]
                for currency, bucket in summary["currencies"].items()
                for key, value in bucket.items()
            ),
            (["", key, value] for key, value in summary["counts"].items()),
            [
                ["", "Importación completa", summary["import_complete"]],
                ["", "Definición", summary["economic_definition"]],
                ["", "Procedencia", provenance_note(run)],
                [
                    "",
                    "Nota",
                    "Importes guardados como texto decimal exacto para evitar pérdida de precisión en Excel.",
                ],
            ],
        ),
    )
    findings = run["result"]["findings"]
    sheet(
        "Hallazgos",
        (
            "Hallazgo",
            "Operaciones",
            "Cargos",
            "Concepto",
            "Moneda",
            "Facturado",
            "Esperado",
            "Diferencia calculada",
            "Diferencia confirmada",
            "Estado",
            "Acuerdo",
            "Versión",
            "Regla",
            "Motivos",
            "Evidencia",
            "Faltante",
        ),
        (
            [
                f["id"],
                ", ".join(f["shipment_ids"]),
                ", ".join(f["charge_ids"]),
                f["concept"],
                f["currency"],
                f["actual"],
                f["expected"],
                f["difference"],
                f["confirmed_difference"],
                STATUS_LABELS[f["status"]],
                f["agreement"],
                f["version"],
                f["rule"],
                "\n".join(f["reasons"]),
                ", ".join(f["evidence_ids"]),
                ", ".join(f["missing_evidence"]),
            ]
            for f in findings
        ),
    )
    sheet(
        "Cálculos",
        ["Hallazgo", "Paso", "Explicación del cálculo", "Resultado", "Datos de ejecución"],
        ([f["id"], *row] for f in findings for trace in f["trace"] for row in trace_rows(trace)),
    )
    sheet(
        "Problemas de datos",
        ["Tipo", "Documento", "Fila", "Campo", "Valor original", "Problema"],
        (
            [i["category"], i["document"], i["row"], i["field"], i["raw"], i["message"]]
            for i in run["result"]["issues"]
        ),
    )
    sheet(
        "Evidencia",
        ["Referencia", "Tipo", "Operaciones", "Cargos", "Documento conservado", "Descripción"],
        (
            [
                e["id"],
                e["kind"],
                ", ".join(e["shipment_ids"]),
                ", ".join(e["charge_ids"]),
                e["document_hash"],
                e["note"],
            ]
            for e in run["snapshot"]["evidence"]
        ),
    )
    sheet(
        "Decisiones",
        (
            "Secuencia",
            "Hallazgo",
            "Decisión",
            "Persona",
            "Motivo",
            "Fecha",
            "Evidencia",
            "Ya conocida",
            "Minutos",
            "Hash",
        ),
        (
            [
                d["seq"],
                d["payload"]["finding_id"],
                d["payload"]["action"],
                d["payload"]["actor"],
                d["payload"]["note"],
                d["created_at"],
                ", ".join(d["payload"]["evidence_ids"]),
                d["payload"]["known_to_client"],
                d["payload"]["review_minutes"],
                d["hash"],
            ]
            for d in run["decisions"]
        ),
    )
    sheet(
        "Origen de datos",
        [
            "Registro",
            "Campo",
            "Archivo",
            "Hoja",
            "Fila",
            "Columna",
            "Valor original",
            "Transformación",
            "Hash del archivo",
        ],
        (
            [
                record["id"],
                field,
                ref["filename"],
                ref["sheet"],
                ref["row"],
                ref["column"],
                ref["raw"],
                ref["transform"],
                ref["document"],
            ]
            for entity in ("shipments", "charges")
            for record in run["snapshot"][entity]
            for field, ref in record["provenance"].items()
        ),
    )
    sheet(
        "Metadata",
        ["Dato", "Valor"],
        [[key, value] for key, value in run.items() if key not in {"result", "snapshot", "decisions"}]
        + [
            ["Motor", run["result"]["engine_version"]],
            ["Datos normalizados", run["result"]["semantic_hash"]],
        ],
    )


def workbook_bytes(run: dict) -> bytes:
    data = io.BytesIO()
    write_workbook(run, data)
    return data.getvalue()


def html_report(run: dict) -> str:
    def escape(value):
        return html.escape(str(value if value is not None else "—"))

    summary = run["result"]["summary"]
    totals = "".join(
        f"<tr><td>{escape(currency)}</td><td>{escape(bucket['actual'])}</td><td>{escape(bucket['determinable'])}</td><td>{escape(bucket['confirmed_overcharge'])}</td><td>{escape(bucket['confirmed_undercharge'])}</td><td>{escape(bucket['review'])}</td><td>{escape(bucket['undeterminable'])}</td></tr>"
        for currency, bucket in summary["currencies"].items()
    )
    rows = "".join(
        f"<tr><td>{escape(', '.join(f['shipment_ids']))}</td><td>{escape(f['concept'])}</td><td>{escape(STATUS_LABELS[f['status']])}</td><td>{escape(f['currency'])} {escape(f['actual'])}</td><td>{escape(f['expected'])}</td><td>{escape(f['confirmed_difference'])}</td><td>{escape(' '.join(f['reasons']))}</td></tr>"
        for f in run["result"]["findings"]
    )
    warnings = (
        "<p class='warning'>IMPORTACIÓN INCOMPLETA: hay filas rechazadas. Los totales sólo incluyen cargos aceptados. Corregir los datos antes de concluir.</p>"
        if not summary["import_complete"]
        else ""
    )
    warnings += "<p>Origen de los datos: " + escape(provenance_note(run)) + "</p>"
    return f"""<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Freight Audit · Informe</title><style>body{{font:14px system-ui;color:#18312e;margin:40px;line-height:1.5}}h1{{font-size:32px}}table{{border-collapse:collapse;width:100%;margin:24px 0;font-size:12px}}th,td{{border-bottom:1px solid #cdd9d4;padding:10px;text-align:left;vertical-align:top}}th{{background:#edf3ef}}.warning{{padding:16px;background:#fff0d5}}small{{overflow-wrap:anywhere}}@media print{{body{{margin:12mm}}thead{{display:table-header-group}}tr{{break-inside:avoid}}}}</style><body><p>FREIGHT AUDIT / INFORME LOCAL</p><h1>{escape(run["snapshot"]["label"])}</h1><p>{summary["determinable_findings"]} de {summary["total_findings"]} hallazgos determinables. Estado técnico y decisión humana se conservan por separado.</p>{warnings}<table><thead><tr><th>Moneda</th><th>Facturado aceptado</th><th>Determinable</th><th>Exceso determinado</th><th>Defecto determinado</th><th>En revisión</th><th>Indeterminado</th></tr></thead><tbody>{totals}</tbody></table><p>{escape(summary["economic_definition"])}</p><table><thead><tr><th>Operaciones</th><th>Concepto</th><th>Estado</th><th>Facturado</th><th>Esperado</th><th>Diferencia confirmada</th><th>Motivo</th></tr></thead><tbody>{rows}</tbody></table><p>Decisiones humanas registradas: {len(run["decisions"])}. El detalle, las fuentes y los cálculos están en el XLSX y en audit.json.</p><small>Auditoría {escape(run["id"])}<br>Resultado {escape(run["result_hash"])}<br>Motor {escape(run["artifact_hash"])}</small></body></html>"""


def _file_chunks(source):
    while data := source.read(1024 * 1024):
        yield data


def write_bundle(store: Store, run_id: str, output, format: str = "freight-audit-bundle/v2") -> None:
    """Write a bounded portable package to a caller-owned stream."""
    if format not in {"freight-audit-bundle/v1", "freight-audit-bundle/v2"}:
        raise ValueError(f"Formato de paquete no admitido: {format}")
    run = store.load(run_id)
    originals = store.source_hashes(Dataset.model_validate(run["snapshot"]))
    with store.connect() as connection:
        source_bytes = 0
        for source in originals:
            row = connection.execute("SELECT length(data) FROM sources WHERE hash=?", (source,)).fetchone()
            if row is None:
                raise IntegrityError("Falta un original del paquete; restaurar una copia verificada.")
            source_bytes += row[0]
            if source_bytes > MAX_BUNDLE_UNCOMPRESSED_BYTES:
                raise ReportLimitError("Los originales del paquete superan el límite admitido sin comprimir.")
    used = 0
    manifest = {}
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:

        def member(name, chunks, *, record=True):
            nonlocal used
            checksum = hashlib.sha256()
            with archive.open(name, "w", force_zip64=True) as destination:
                for chunk in chunks:
                    used += len(chunk)
                    if used > MAX_BUNDLE_UNCOMPRESSED_BYTES:
                        raise ReportLimitError("El paquete portable supera el límite admitido sin comprimir.")
                    checksum.update(chunk)
                    destination.write(chunk)
            if record:
                manifest[name] = checksum.hexdigest()

        member("audit.json", canonical_chunks(run))
        if format == "freight-audit-bundle/v1":
            member("snapshot.json", canonical_chunks(run["snapshot"]))
        member("reporte.html", (html_report(run).encode(),))
        # Reject a doomed package before allocating/generating its workbook.
        if used + source_bytes > MAX_BUNDLE_UNCOMPRESSED_BYTES:
            raise ReportLimitError("El paquete portable supera el límite admitido sin comprimir.")
        with tempfile.TemporaryFile(mode="w+b") as spreadsheet:
            try:
                write_workbook(run, spreadsheet)
            except ReportLimitError as exc:
                member(
                    "ADVERTENCIA_EXPORTACION.txt",
                    ((str(exc) + "\nNo se incluyó una planilla parcial.\n").encode(),),
                )
            else:
                if spreadsheet.tell() + used + source_bytes > MAX_BUNDLE_UNCOMPRESSED_BYTES:
                    raise ReportLimitError(
                        "La planilla y los originales superan el límite del paquete; dividir el período o exportar sus datos por separado."
                    )
                spreadsheet.seek(0)
                member("auditoria.xlsx", _file_chunks(spreadsheet))
        for source_hash in originals:
            # One original at a time; never retain all documents in a dict.
            member(f"sources/{source_hash}", (store.source(source_hash),))
        member(
            "manifest.json",
            canonical_chunks({"format": format, "run_id": run_id, "files": manifest}),
            record=False,
        )


def bundle_bytes(store: Store, run_id: str, format: str = "freight-audit-bundle/v2") -> bytes:
    """Compatibility adapter for consumers explicitly requesting bytes."""
    with tempfile.SpooledTemporaryFile(max_size=1024 * 1024, mode="w+b") as output:
        write_bundle(store, run_id, output, format)
        output.seek(0)
        return output.read()


def verify_bundle(data) -> dict:
    try:
        if isinstance(data, (str, Path)):
            details = Path(data).stat()
            if not stat.S_ISREG(details.st_mode) or details.st_size > MAX_BUNDLE_UNCOMPRESSED_BYTES:
                raise IntegrityError("Seleccionar un paquete regular dentro del límite de tamaño admitido.")
        with zipfile.ZipFile(io.BytesIO(data) if isinstance(data, bytes) else data) as archive:
            infos = archive.infolist()
            if (
                len({i.filename for i in infos}) != len(infos)
                or any(i.flag_bits & 1 for i in infos)
                or sum(i.file_size for i in infos) > MAX_BUNDLE_UNCOMPRESSED_BYTES
            ):
                raise IntegrityError("El paquete contiene entradas repetidas o supera el tamaño admitido.")
            manifest = load_json(archive.read("manifest.json"))
            if not isinstance(manifest, dict) or not isinstance(manifest.get("files"), dict):
                raise IntegrityError("El manifiesto del paquete no tiene la estructura requerida.")
            bundle_format = manifest.get("format")
            if bundle_format not in {"freight-audit-bundle/v1", "freight-audit-bundle/v2"}:
                raise IntegrityError("Formato de paquete no admitido.")
            if set(archive.namelist()) != {*manifest["files"], "manifest.json"}:
                raise IntegrityError("El paquete contiene archivos no declarados o faltantes.")
            for name, source_hash in manifest["files"].items():
                if (
                    name.startswith("/")
                    or "\\" in name
                    or ":" in name
                    or any(part in {"", ".", ".."} for part in name.split("/"))
                    or ".." in Path(name).parts
                    or _archive_hash(archive, name) != source_hash
                ):
                    raise IntegrityError("Un archivo del paquete fue alterado.")
            if "reporte.html" not in manifest["files"] or not (
                {"auditoria.xlsx", "ADVERTENCIA_EXPORTACION.txt"} & manifest["files"].keys()
            ):
                raise IntegrityError("Falta un informe o la advertencia de exportación del paquete.")
            run = load_json(archive.read("audit.json"))
            if (
                digest(run["snapshot"]) != run["input_hash"]
                or digest(run["result"]) != run["result_hash"]
                or run_identity(run["input_hash"], run["result_hash"], run["artifact_hash"], run["metadata"])
                != run["id"]
                or run["id"] != manifest["run_id"]
            ):
                raise IntegrityError("Los hashes de auditoría no coinciden.")
            if bundle_format == "freight-audit-bundle/v1":
                if "snapshot.json" not in archive.namelist():
                    raise IntegrityError("El paquete v1 no contiene snapshot.json.")
                if canonical(load_json(archive.read("snapshot.json"))) != canonical(run["snapshot"]):
                    raise IntegrityError("El snapshot no coincide con la auditoría.")
            elif bundle_format == "freight-audit-bundle/v2":
                if "snapshot.json" in archive.namelist():
                    raise IntegrityError("El paquete v2 no admite snapshot.json redundante.")
            previous = run["id"]
            for decision in run["decisions"]:
                if (
                    decision["previous_hash"] != previous
                    or digest([previous, decision["payload"], decision["created_at"]]) != decision["hash"]
                ):
                    raise IntegrityError("La cadena de decisiones no coincide.")
                previous = decision["hash"]
            for source_hash in Store.source_hashes(Dataset.model_validate(run["snapshot"])):
                if manifest["files"].get(f"sources/{source_hash}") != source_hash:
                    raise IntegrityError("El documento original fue alterado.")
            return run
    except (KeyError, ValueError, TypeError, AttributeError, zipfile.BadZipFile) as exc:
        if isinstance(exc, IntegrityError):
            raise
        raise IntegrityError("El paquete está incompleto o dañado; usar una copia verificada.") from exc


def _archive_hash(archive, name):
    checksum = hashlib.sha256()
    with archive.open(name) as source:
        for chunk in _file_chunks(source):
            checksum.update(chunk)
    return checksum.hexdigest()


def export_run(store: Store, run_id: str, destination: Path) -> dict:
    destination = checked_output_path(destination)
    if destination.exists():
        raise ValueError("La carpeta de salida ya existe; elegir una carpeta nueva.")
    with tempfile.TemporaryFile(mode="w+b") as bundle:
        write_bundle(store, run_id, bundle)
        bundle.seek(0)
        destination.parent.mkdir(parents=True, exist_ok=True)
        checked_output_path(destination)
        destination.mkdir(exist_ok=False)
        incomplete = destination / "EXPORTACION_INCOMPLETA.txt"
        incomplete.write_text(
            "La exportación se interrumpió. No usar estos archivos; repetir en una carpeta nueva.\n",
            encoding="utf-8",
        )
        with zipfile.ZipFile(bundle) as archive:
            for name in archive.namelist():
                if name not in {
                    "audit.json",
                    "reporte.html",
                    "auditoria.xlsx",
                    "ADVERTENCIA_EXPORTACION.txt",
                    "snapshot.json",
                    "manifest.json",
                } and not re.fullmatch(r"sources/[0-9a-f]{64}", name):
                    raise IntegrityError("El paquete contiene una ruta de salida no permitida.")
                target = destination / name
                if target.parent != destination and not target.parent.exists():
                    target.parent.mkdir(exist_ok=False)
                checked_output_path(target)
                with target.open("xb") as output, archive.open(name) as source:
                    shutil.copyfileobj(source, output, length=1024 * 1024)
        checksum = hashlib.sha256()
        bundle.seek(0)
        with (destination / "auditoria.zip").open("xb") as output:
            for chunk in _file_chunks(bundle):
                output.write(chunk)
                checksum.update(chunk)
        incomplete.unlink()
        return {"directory": str(destination), "run_id": run_id, "bundle_hash": checksum.hexdigest()}
