"""Entirely fictional agreements. This module is outside the generic calculation core."""

import csv
import json
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

from .engine import audit
from .storage import Store


def const(value, unit=None, kind="decimal"):
    return {"op": "const", "value": {"type": kind, "value": value, "unit": unit}}


def attr(field, unit=None):
    return {"op": "attr", "field": field, **({"unit": unit} if unit else {})}


def op(name, *args, **extra):
    return {"op": name, "args": list(args), **extra}


def rule(identifier, concept, expression, **extra):
    return {
        "id": identifier,
        "concept": concept,
        "description": f"Regla ficticia: {concept}",
        "expression": expression,
        **extra,
    }


def agreement(identifier, carrier, versions, **extra):
    return {
        "id": identifier,
        "name": f"Acuerdo FICTICIO {identifier}",
        "carrier": carrier,
        "currency": "ARS",
        "date_field": "service_date",
        "scale": 2,
        "rounding": "ROUND_HALF_UP",
        "tolerance_absolute": "0.01",
        "tolerance_relative": "0",
        "matching": {
            "keys": [{"charge": "reference", "shipment": "reference"}],
            "cardinality": "one",
            "duplicate_fields": ["reference", "concept", "amount"],
        },
        "versions": versions,
        **extra,
    }


def version(identifier, rules, start="2026-01-01", end=None, tables=None):
    return {
        "id": identifier,
        "valid_from": start,
        "valid_to": end,
        "rules": rules,
        "tables": tables or {},
        "source_note": "Ejemplo totalmente ficticio. Sin valor contractual ni evidencia de prácticas del mercado.",
    }


def column(source, target, kind="text", **extra):
    return {"source": source, "target": target, "type": kind, **extra}


def constant(target, value):
    return {"target": target, "type": "text", "constant": value}


def ar_number(value):
    return str(value).replace(".", ",")


def generate(root: Path):
    root.mkdir(parents=True, exist_ok=True)
    base_a = op(
        "mul",
        op("max", op("mul", attr("weight", "kg"), const("2", "ARS/kg")), const("300", "ARS")),
        const("1.05"),
    )
    rules_a = [
        rule("A-BASE", "base", base_a, expected=True),
        rule(
            "A-WAIT",
            "wait",
            const("80", "ARS"),
            evidence=[{"any_of": ["authorization", "signed_event"], "scope": "each_shipment"}],
        ),
    ]
    a = agreement(
        "A",
        "FICTICIO-ALFA",
        [
            version("A-2026-H1", rules_a, end="2026-06-30"),
            version(
                "A-2026-H2",
                [rule("A-BASE-2", "base", op("mul", base_a, const("1.1")), expected=True), rules_a[1]],
                start="2026-07-01",
            ),
            version("A-SOLAPADA", rules_a, start="2026-07-15", end="2026-07-20"),
        ],
        detect_missing=True,
    )
    b_table = {
        "zones": {
            "kind": "lookup",
            "rows": [
                {"keys": [zone, vehicle], "value": {"type": "decimal", "value": amount, "unit": "ARS"}}
                for zone, vehicle, amount in [
                    ("Norte", "Furgón", "900"),
                    ("Sur", "Furgón", "1100"),
                    ("Norte", "Camión", "1500"),
                    ("Sur", "Camión", "1700"),
                ]
            ],
        }
    }
    b = agreement(
        "B",
        "FICTICIO-BETA",
        [
            version(
                "B-1",
                [rule("B-ZONA", "base", op("lookup", attr("zone"), attr("vehicle"), table="zones"))],
                tables=b_table,
            )
        ],
    )
    c_table = {
        "weights": {
            "kind": "band",
            "unit": "kg",
            "bands": [
                {"lower": "0", "upper": "500", "value": {"type": "decimal", "value": "400", "unit": "ARS"}},
                {
                    "lower": "500",
                    "upper": "1500",
                    "value": {"type": "decimal", "value": "700", "unit": "ARS"},
                },
                {"lower": "1500", "value": {"type": "decimal", "value": "1000", "unit": "ARS"}},
            ],
        }
    }
    c = agreement(
        "C",
        "FICTICIO-GAMMA",
        [
            version(
                "C-1",
                [
                    rule("C-BANDA", "base", op("band", attr("weight", "kg"), table="weights")),
                    rule(
                        "C-PALLET", "pallet", op("mul", attr("pallets", "pallet"), const("30", "ARS/pallet"))
                    ),
                    rule(
                        "C-EXTRA",
                        "special",
                        const("120", "ARS"),
                        evidence=[{"any_of": ["authorization"], "document_required": True}],
                    ),
                ],
                tables=c_table,
            )
        ],
    )
    agreements = [a, b, c]
    shipment_rows = []
    charge_rows: list[list] = []
    evidence: list[dict] = []

    def charge(ref, carrier, agreement_id, concept, amount, suffix=""):
        charge_rows.append(
            [
                f"C-{len(charge_rows) + 1:03d}{suffix}",
                "L-FICTICIA-01",
                carrier,
                agreement_id,
                ref,
                concept,
                ar_number(amount),
                "ARS",
            ]
        )

    for i in range(1, 33):
        group = "A" if i <= 16 else "B" if i <= 24 else "C"
        carrier = {"A": "FICTICIO-ALFA", "B": "FICTICIO-BETA", "C": "FICTICIO-GAMMA"}[group]
        ref = f"{i:05d}"
        weight = 100 + i * 80
        zone = "Norte" if i % 2 else "Sur"
        vehicle = "Camión" if i % 3 == 0 else "Furgón"
        service_date = "15/05/2026" if i != 5 else "15/05/2025"
        if i == 6:
            service_date = "16/07/2026"
        if i == 13:
            service_date = "01/08/2026"
        shipment_rows.append(
            [
                f"S-{i:03d}",
                ref,
                carrier,
                service_date,
                "" if i == 4 else ar_number(weight),
                zone,
                vehicle,
                str(i % 4 + 1),
                "Santa Fe",
                "Rosario",
            ]
        )
        if group == "A":
            expected: int | Decimal = max(weight * 2, 300) * 105 // 100
            if i == 13:
                expected = Decimal(expected) * Decimal("1.1")
            if i == 7:
                charge(ref, carrier, group, "Concepto sin clasificar", 432)
            elif i == 9:
                charge(ref, carrier, group, "Flete", 100)
                charge(ref, carrier, group, "Flete", expected - 100)
            elif i != 12:
                charge(ref, carrier, group, "Flete", expected + (50 if i == 2 else 0))
            if i == 8:
                charge(ref, carrier, group, "Flete", expected)
            if i in {3, 10}:
                charge(ref, carrier, group, "Demora", 80)
                if i == 10:
                    evidence.append(
                        {
                            "id": "EV-10",
                            "kind": "signed_event",
                            "shipment_ids": [f"S-{i:03d}"],
                            "note": "Declaración ficticia de autorización; sin documento adjunto requerido en este ejemplo.",
                        }
                    )
        elif group == "B":
            expected = {
                ("Norte", "Furgón"): 900,
                ("Sur", "Furgón"): 1100,
                ("Norte", "Camión"): 1500,
                ("Sur", "Camión"): 1700,
            }[(zone, vehicle)]
            charge(ref, carrier, group, "Flete", expected - (75 if i == 18 else 0))
        else:
            expected = 400 if weight < 500 else 700 if weight < 1500 else 1000
            charge(ref, carrier, group, "Flete", expected)
            charge(ref, carrier, group, "Pallet", (i % 4 + 1) * 30)
            if i in {25, 26}:
                charge(ref, carrier, group, "Especial", 120)
                if i == 26:
                    evidence.append(
                        {
                            "id": "EV-26",
                            "kind": "authorization",
                            "shipment_ids": [f"S-{i:03d}"],
                            "note": "Autorización ficticia en documento de demostración.",
                            "document_hash": None,
                        }
                    )
    charge("INEXISTENTE", "FICTICIO-ALFA", "A", "Flete", 123)
    workbook = Workbook()
    info = workbook.active
    info.title = "LEER"
    info.append(["DATOS TOTALMENTE FICTICIOS. No representan tarifas ni usos del mercado."])
    sheet = workbook.create_sheet("Despachos")
    sheet.append(["Exportación de ejemplo — FICTICIO"])
    headers = [
        "Registro",
        "Numero Rem.",
        "Empresa transporte",
        "Fecha Despacho",
        "Peso KG",
        "Zona entrega",
        "Vehículo",
        "Pallets",
        "Origen",
        "Destino",
    ]
    sheet.append(headers)
    for row in shipment_rows:
        sheet.append(row)
    workbook.save(root / "operaciones.xlsx")
    workbook.close()
    with (root / "liquidacion.csv").open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file, delimiter=";")
        writer.writerow(
            ["Línea", "Liquidación", "Transportista", "Contrato", "Remito", "Servicio", "Importe", "Moneda"]
        )
        writer.writerows(charge_rows)
    mapping_s = {
        "id": "operaciones-demo",
        "version": "1",
        "entity": "shipments",
        "sheet": "Despachos",
        "header_row": 2,
        "columns": [
            column("Registro", "id"),
            column("Numero Rem.", "reference"),
            column("Empresa transporte", "carrier"),
            column("Fecha Despacho", "attributes.service_date", "date"),
            column("Peso KG", "attributes.weight", "decimal", unit="kg", required=False),
            column("Zona entrega", "attributes.zone"),
            column("Vehículo", "attributes.vehicle"),
            column("Pallets", "attributes.pallets", "decimal", unit="pallet"),
            column("Origen", "attributes.origin"),
            column("Destino", "attributes.destination"),
        ],
    }
    mapping_c = {
        "id": "cargos-demo",
        "version": "1",
        "entity": "charges",
        "columns": [
            column(source, target, "decimal" if target == "amount" else "text")
            for source, target in [
                ("Línea", "id"),
                ("Liquidación", "settlement"),
                ("Transportista", "carrier"),
                ("Contrato", "agreement"),
                ("Remito", "reference"),
                ("Servicio", "concept"),
                ("Importe", "amount"),
                ("Moneda", "currency"),
            ]
        ],
        "concept_map": {"Flete": "base", "Demora": "wait", "Pallet": "pallet", "Especial": "special"},
    }
    for name, data in [
        ("agreements.json", agreements),
        ("mapping-shipments.json", mapping_s),
        ("mapping-charges.json", mapping_c),
        ("evidence.json", evidence),
    ]:
        (root / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (root / "authorization.txt").write_text(
        "DOCUMENTO FICTICIO. Autorización del servicio especial para S-026. Sólo prueba técnica.\n",
        encoding="utf-8",
    )
    (root / "operaciones-invalidas.csv").write_text(
        "Registro;Numero Rem.;Empresa transporte;Fecha Despacho;Peso KG;Zona entrega;Vehículo;Pallets;Origen;Destino\nERR-1;00099;FICTICIO-ALFA;31/13/2026;1.23,45;Norte;Furgón;1;A;B\n",
        encoding="utf-8",
    )
    mapping_bad = {**mapping_s, "id": "errores-demo", "sheet": None, "header_row": 1}
    (root / "mapping-invalid.json").write_text(
        json.dumps(mapping_bad, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (root / "project.json").write_text(
        json.dumps(
            {
                "label": "Demostración ficticia · 3 acuerdos / 32 operaciones",
                "shipments": {"file": "operaciones.xlsx", "mapping": "mapping-shipments.json"},
                "charges": {"file": "liquidacion.csv", "mapping": "mapping-charges.json"},
                "agreements": "agreements.json",
                "evidence": "evidence.json",
                "attachments": [{"file": "authorization.txt", "evidence_id": "EV-26"}],
                "coverage": {"A": [f"S-{i:03d}" for i in range(1, 17)]},
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    generate_second(root / "second-client")


def generate_second(root: Path):
    root.mkdir(exist_ok=True)
    # Different column names, decimal locale, currency, date field, consolidation and units.
    expr = op(
        "mul",
        op(
            "max",
            {"op": "sum", "field": "mass", "unit": "kg"},
            op("mul", {"op": "sum", "field": "cube", "unit": "m3"}, const("140", "kg/m3")),
        ),
        const("0.37", "USD/kg"),
    )
    a = agreement(
        "D",
        "FICTITIOUS-OMEGA",
        [version("D-1", [rule("D-VOLUME", "movement", expr)])],
        currency="USD",
        date_field="pickup",
        matching={
            "keys": [{"charge": "attributes.batch", "shipment": "attributes.batch"}],
            "cardinality": "group",
        },
    )
    with (root / "movements.csv").open("w", newline="") as file:
        writer = csv.writer(file, delimiter="\t")
        writer.writerow(["Movement key", "Waybill", "Picked up", "Lot", "Mass", "Cube"])
        for i in range(1, 7):
            writer.writerow(
                [f"M{i}", f"WB-{i:05d}", "2026-09-01", "LOT-X" if i < 4 else "LOT-Y", "100.25", "1.25"]
            )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoice detail"
    sheet.append(["Row key", "Lot", "Net USD", "Charge label"])
    sheet.append(["INV-1", "LOT-X", "194.25", "Consolidated movement"])
    sheet.append(["INV-2", "LOT-Y", "204.25", "Consolidated movement"])
    workbook.save(root / "invoice.xlsx")
    workbook.close()
    s = {
        "id": "omega-ops",
        "version": "1",
        "entity": "shipments",
        "delimiter": "\t",
        "decimal_separator": ".",
        "thousands_separator": "",
        "date_formats": ["%Y-%m-%d"],
        "columns": [
            column("Movement key", "id"),
            column("Waybill", "reference"),
            constant("carrier", "FICTITIOUS-OMEGA"),
            column("Picked up", "attributes.pickup", "date"),
            column("Lot", "attributes.batch"),
            column("Mass", "attributes.mass", "decimal", unit="kg"),
            column("Cube", "attributes.cube", "decimal", unit="m3"),
        ],
    }
    c = {
        "id": "omega-invoice",
        "version": "1",
        "entity": "charges",
        "sheet": "Invoice detail",
        "decimal_separator": ".",
        "thousands_separator": "",
        "columns": [
            column("Row key", "id"),
            column("Lot", "reference"),
            column("Lot", "attributes.batch"),
            constant("carrier", "FICTITIOUS-OMEGA"),
            constant("settlement", "INV-EXAMPLE"),
            constant("agreement", "D"),
            constant("currency", "USD"),
            column("Net USD", "amount", "decimal"),
            column("Charge label", "concept"),
        ],
        "concept_map": {"Consolidated movement": "movement"},
    }
    for name, value in [
        ("agreements.json", [a]),
        ("mapping-shipments.json", s),
        ("mapping-charges.json", c),
        (
            "project.json",
            {
                "label": "Segundo cliente FICTICIO · consolidado / USD",
                "shipments": {"file": "movements.csv", "mapping": "mapping-shipments.json"},
                "charges": {"file": "invoice.xlsx", "mapping": "mapping-charges.json"},
                "agreements": "agreements.json",
            },
        ),
    ]:
        (root / name).write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def run_demo(store: Store, root: Path):
    from .project import load_project

    dataset, imports = load_project(root / "project.json", store)
    result = audit(dataset)
    run_id = store.save(dataset, result)
    return {"run_id": run_id, "summary": result.summary, "imports": imports}
