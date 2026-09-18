#!/usr/bin/env python3
"""Minimal Reproducer for PERF-01: Quadratic Complexity O(R^2) in reporting.py sheet()

Demonstrates how accessing `ws.max_row` on every row insertion in openpyxl causes
an O(R^2) performance cliff during workbook_bytes() generation.
"""

import time
from openpyxl import Workbook

def test_quadratic_max_row(row_count: int):
    wb = Workbook()
    ws = wb.active
    ws.append(["col1", "col2", "col3", "col4", "col5"])
    
    t0 = time.perf_counter()
    for i in range(row_count):
        # This is exactly what src/freight_audit/reporting.py line 80 does:
        if ws.max_row >= 1_048_576:
            break
        ws.append([f"val_{i}_1", f"val_{i}_2", f"val_{i}_3", f"val_{i}_4", f"val_{i}_5"])
    duration = time.perf_counter() - t0
    return duration

def test_linear_counter(row_count: int):
    wb = Workbook()
    ws = wb.active
    ws.append(["col1", "col2", "col3", "col4", "col5"])
    
    t0 = time.perf_counter()
    curr_row = 1
    for i in range(row_count):
        if curr_row >= 1_048_576:
            break
        ws.append([f"val_{i}_1", f"val_{i}_2", f"val_{i}_3", f"val_{i}_4", f"val_{i}_5"])
        curr_row += 1
    duration = time.perf_counter() - t0
    return duration

if __name__ == "__main__":
    print("Benchmarking openpyxl row insertion with ws.max_row vs integer counter:")
    for n in [1000, 3000, 10000]:
        t_quad = test_quadratic_max_row(n)
        t_lin = test_linear_counter(n)
        print(f"N={n:5d} rows | with ws.max_row: {t_quad:6.3f}s | with counter: {t_lin:6.3f}s | ratio: {t_quad/max(t_lin, 0.001):.1f}x slower")
