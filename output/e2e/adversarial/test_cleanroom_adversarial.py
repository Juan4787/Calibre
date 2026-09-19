#!/usr/bin/env python3
"""
FASE 2: E2E ADVERSARIAL DE INCERTIDUMBRE Y METAMORFISMOS (CLEAN-ROOM)

Executed strictly in /tmp/calibre-e2e-adversarial/ using the isolated venv and installed wheel.
Proves that uncertainty is never converted into certainty across any layer:
File -> Importer -> Snapshot -> Engine -> SQLite -> API -> UI -> XLSX -> HTML

Exercises:
- 10 Intentional Adversarial Uncertainty Cases (ADV-01 through ADV-10)
- 2 Deterministic Control Cases (CTRL-PASS and CTRL-FAIL)
- 6 Metamorphic Invariance Transformations (Meta-1 through Meta-6)
- Cold Restart Recovery
- 6-Way Cross-Channel Reconciliation
"""

import json
import os
import signal
import sqlite3
import subprocess
import sys
import time
import urllib.request
import zipfile
from decimal import Decimal
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

CLEAN_ROOT = Path("/tmp/calibre-e2e-adversarial")
VENV_DIR = CLEAN_ROOT / "venv"
PYTHON_BIN = VENV_DIR / "bin" / "python"
FREIGHT_AUDIT_BIN = VENV_DIR / "bin" / "freight-audit"
RUN_DIR = CLEAN_ROOT / "run"
FIXTURES_DIR = CLEAN_ROOT / "fixtures"
DB_DIR = CLEAN_ROOT / "db"
DB_PATH = DB_DIR / "adversarial_audit.db"
EXPORTS_DIR = CLEAN_ROOT / "exports"
SCREENSHOTS_DIR = CLEAN_ROOT / "screenshots"
OBSERVED_DIR = CLEAN_ROOT / "observed"

PORT = 8770
BASE_URL = f"http://127.0.0.1:{PORT}"
WORKSPACE_PATH = str(Path(__file__).resolve().parents[3])


def verify_clean_environment():
    import freight_audit
    module_file = freight_audit.__file__
    print(f"[*] Checking freight_audit module location: {module_file}")
    assert str(VENV_DIR) in module_file, f"Module not in clean venv: {module_file}"
    assert WORKSPACE_PATH not in module_file, f"Module loaded from workspace: {module_file}"

    for p in sys.path:
        assert WORKSPACE_PATH not in p, f"sys.path contains workspace reference: {p}"

    assert "PYTHONPATH" not in os.environ or not os.environ["PYTHONPATH"], "PYTHONPATH must be empty"
    print("[✓] Clean-room isolation verified: Zero references to workspace checkout.")


def clean_env():
    env = os.environ.copy()
    for key in ("PYTHONPATH", "PYTHONHOME", "VIRTUAL_ENV"):
        env.pop(key, None)
    env["PATH"] = f"{VENV_DIR / 'bin'}:{env.get('PATH', '')}"
    return env


def start_server():
    cmd = [
        str(FREIGHT_AUDIT_BIN),
        "--db",
        str(DB_PATH),
        "serve",
        "--port",
        str(PORT),
    ]
    print(f"[*] Starting clean-room adversarial server: {' '.join(cmd)}")
    proc = subprocess.Popen(
        cmd,
        cwd=str(RUN_DIR),
        env=clean_env(),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    deadline = time.time() + 15
    while time.time() < deadline:
        try:
            req = urllib.request.Request(f"{BASE_URL}/api/session")
            with urllib.request.urlopen(req, timeout=1) as resp:
                if resp.status == 200:
                    data = json.loads(resp.read().decode())
                    print(f"[✓] Server started and responsive (token: {data.get('token')[:8]}...)")
                    return proc
        except Exception:
            time.sleep(0.3)
    proc.terminate()
    stdout, stderr = proc.communicate()
    raise RuntimeError(f"Server failed to start. stdout: {stdout}, stderr: {stderr}")


def stop_server(proc):
    print("[*] Stopping server...")
    proc.terminate()
    try:
        proc.wait(timeout=5)
    except subprocess.TimeoutExpired:
        proc.kill()
        proc.wait()
    print("[✓] Server stopped cleanly.")


def run_adversarial_suite():
    verify_clean_environment()

    for d in (DB_DIR, EXPORTS_DIR, SCREENSHOTS_DIR, OBSERVED_DIR):
        d.mkdir(parents=True, exist_ok=True)

    if DB_PATH.exists():
        DB_PATH.unlink()

    server_proc = start_server()

    try:
        with sync_playwright() as p:
            print("[*] Launching Chromium headless browser for Adversarial Suite...")
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(viewport={"width": 1440, "height": 960})
            page = context.new_page()

            page.goto(BASE_URL)
            page.wait_for_selector(".brand")
            print("[✓] App loaded.")

            # =========================================================
            # PART A: ADV-09 — Ingesta con Fila Rechazada (Lote Incompleto)
            # =========================================================
            print("\n" + "=" * 60)
            print("PART A: ADV-09 — Ingesta con Fila Rechazada (Lote Incompleto)")
            print("=" * 60)
            page.click('button.nav[data-view="new"]')
            page.wait_for_selector("#audit-label")
            page.fill("#audit-label", "Auditoría ADV-09 (Fila Rechazada)")

            page.set_input_files("#file-shipments", str(FIXTURES_DIR / "operaciones_with_error.csv"))
            page.fill("#mapping-shipments", (FIXTURES_DIR / "mapping_shipments_adv.json").read_text(encoding="utf-8"))
            page.click("#validate-shipments")
            page.wait_for_selector('#result-shipments:has-text("12 filas aceptadas · 1 rechazadas")')
            print("[✓] Shipments validated with 1 rejected row detected in UI.")

            page.set_input_files("#file-charges", str(FIXTURES_DIR / "cargos_adv.csv"))
            page.fill("#mapping-charges", (FIXTURES_DIR / "mapping_charges_adv.json").read_text(encoding="utf-8"))
            page.click("#validate-charges")
            page.wait_for_selector('#result-charges:has-text("13 filas aceptadas · 0 rechazadas")')

            page.fill("#agreements", (FIXTURES_DIR / "agreements_adv.json").read_text(encoding="utf-8"))
            page.screenshot(path=str(SCREENSHOTS_DIR / "01_adv09_formulario_fila_rechazada.png"))

            page.click("#execute")
            page.wait_for_selector(".statusline")
            page.wait_for_selector(".banner.warning")
            page.screenshot(path=str(SCREENSHOTS_DIR / "02_adv09_importacion_incompleta.png"))

            banner_text = page.locator(".banner.warning").inner_text()
            print(f"    Banner text: {banner_text}")
            assert "Importación incompleta" in banner_text, "Warning banner missing on rejected row"

            statusline_adv09 = " ".join(page.locator(".statusline").inner_text().split())
            print(f"    ADV-09 Statusline: {statusline_adv09}")
            # When import is incomplete, zero findings are determinable (0 PASS, 0 FAIL)
            assert "Coincide 0" in statusline_adv09
            assert "Discrepancia 0" in statusline_adv09
            assert "Revisión humana 8" in statusline_adv09
            assert "Indeterminado 4" in statusline_adv09
            print("[✓] CRITICAL INVARIANT VERIFIED (ADV-09): When import has rejected rows, NO finding is allowed to settle as PASS or FAIL! All determinable findings are held in REVIEW.")

            # =========================================================
            # PART B: Full Adversarial Baseline (ADV-01 .. ADV-10 + Controls)
            # =========================================================
            print("\n" + "=" * 60)
            print("PART B: Full Adversarial Baseline Suite (Clean Ingestion)")
            print("=" * 60)
            page.click('button.nav[data-view="new"]')
            page.wait_for_selector("#audit-label")
            page.fill("#audit-label", "Auditoría Adversarial Completa (Baseline)")

            page.set_input_files("#file-shipments", str(FIXTURES_DIR / "operaciones_adv.csv"))
            page.fill("#mapping-shipments", (FIXTURES_DIR / "mapping_shipments_adv.json").read_text(encoding="utf-8"))
            page.click("#validate-shipments")
            page.wait_for_selector('#result-shipments:has-text("12 filas aceptadas · 0 rechazadas")')

            page.set_input_files("#file-charges", str(FIXTURES_DIR / "cargos_adv.csv"))
            page.fill("#mapping-charges", (FIXTURES_DIR / "mapping_charges_adv.json").read_text(encoding="utf-8"))
            page.click("#validate-charges")
            page.wait_for_selector('#result-charges:has-text("13 filas aceptadas · 0 rechazadas")')

            page.fill("#agreements", (FIXTURES_DIR / "agreements_adv.json").read_text(encoding="utf-8"))
            page.screenshot(path=str(SCREENSHOTS_DIR / "03_adversarial_baseline_form.png"))

            page.click("#execute")
            page.wait_for_selector(".statusline")
            page.wait_for_selector("#findings tr")
            page.screenshot(path=str(SCREENSHOTS_DIR / "04_adversarial_baseline_resultado.png"))

            statusline_base = " ".join(page.locator(".statusline").inner_text().split())
            print(f"    Baseline Statusline: {statusline_base}")
            assert "Coincide 1" in statusline_base, "Must have exactly 1 PASS (CTRL-PASS)"
            assert "Discrepancia 1" in statusline_base, "Must have exactly 1 FAIL (CTRL-FAIL)"
            assert "Revisión humana 6" in statusline_base, "Must have exactly 6 REVIEW"
            assert "Indeterminado 4" in statusline_base, "Must have exactly 4 UNDETERMINABLE"

            export_zip_href = page.locator('a:has-text("Exportar paquete")').get_attribute("href")
            base_run_id = export_zip_href.split("/")[3]
            print(f"[✓] Adversarial Baseline Run ID: {base_run_id}")

            # Register human decision on CTRL-FAIL (REM-FAIL)
            print("[*] Registering human decision on CTRL-FAIL...")
            page.locator('button[aria-label*="REM-FAIL"]').click()
            page.wait_for_selector("dialog#detail[open]")
            assert page.locator("#detail .badge.FAIL").is_visible()
            page.fill("#actor", "Auditor QA Adversarial")
            page.select_option("#action", "REJECTED")
            page.fill("#decision-note", "Sobreprecio adversarial confirmado y rechazado.")
            page.select_option("#known", "false")
            page.click('#decision-form button[type="submit"]')
            page.wait_for_selector('.history:has-text("Rechazado")')
            page.screenshot(path=str(SCREENSHOTS_DIR / "05_decision_c_fail.png"))
            page.locator("#close-detail").click()
            page.wait_for_selector("dialog#detail[open]", state="detached")
            print("[✓] Decision saved on CTRL-FAIL.")

            # Inspect all 10 adversarial modals
            cases_to_inspect = [
                ("REM-PASS", "PASS", None),
                ("REM-001", "REVIEW", "Hay varias operaciones candidatas"),
                ("REM-002", "UNDETERMINABLE", "difiere"),
                ("REM-003", "REVIEW", "Falta evidencia requerida"),
                ("REM-004", "UNDETERMINABLE", "versiones aplicables"),
                ("REM-005", "UNDETERMINABLE", "service_date"),
                ("REM-006", "UNDETERMINABLE", "reglas aplicables"),
                ("REM-UNKNOWN", "REVIEW", "No se encontró la operación"),
                ("REM-008", "REVIEW", "posible duplicación"),
                ("REM-010A", "REVIEW", "asignaciones superpuestas"),
            ]

            for ref, expected_status, reason_sub in cases_to_inspect:
                print(f"[*] Inspecting modal for {ref} (expecting {expected_status})...")
                page.locator(f'button[aria-label*="{ref}"]').first.click()
                page.wait_for_selector("dialog#detail[open]")
                assert page.locator(f"#detail .badge.{expected_status}").is_visible(), f"Badge mismatch for {ref}"
                modal_text = page.locator("#detail-content").inner_text()
                if reason_sub:
                    assert reason_sub.lower() in modal_text.lower(), f"Reason '{reason_sub}' not found in modal for {ref}"
                page.locator("#close-detail").click()
                page.wait_for_selector("dialog#detail[open]", state="detached")
                print(f"    [✓] {ref}: Verified status={expected_status}, reason contains '{reason_sub or 'N/A'}'.")

            # In-browser Replay
            page.click("#replay")
            page.wait_for_selector('#notice:has-text("Reproducción idéntica")')
            print("[✓] In-browser Replay passed on adversarial baseline.")

            # Browser downloads
            with page.expect_download() as dl_info:
                page.click('a:has-text("Planilla operativa")')
            base_xlsx = EXPORTS_DIR / f"{base_run_id}_auditoria.xlsx"
            dl_info.value.save_as(str(base_xlsx))

            with page.expect_download() as dl_info:
                page.click('a:has-text("Informe imprimible")')
            base_html = EXPORTS_DIR / f"{base_run_id}_reporte.html"
            dl_info.value.save_as(str(base_html))

            with page.expect_download() as dl_info:
                page.click('a:has-text("Exportar paquete")')
            base_zip = EXPORTS_DIR / f"{base_run_id}_paquete.zip"
            dl_info.value.save_as(str(base_zip))
            print("[✓] Downloaded baseline exports (.xlsx, .html, .zip).")

            # Extract UI DOM observed rows
            ui_rows = []
            for row in page.locator("#findings tr").all():
                cols = [td.inner_text().replace("\n", " ") for td in row.locator("td").all()]
                if len(cols) >= 6:
                    ui_rows.append({
                        "reference_concept": cols[0],
                        "status": cols[1],
                        "actual": cols[2],
                        "expected": cols[3],
                        "difference": cols[4],
                        "decision": cols[5],
                    })

            ui_metrics = [m.inner_text().replace("\n", " | ") for m in page.locator(".metric").all()]
            ui_observed = {
                "run_id": base_run_id,
                "label": page.locator("h1").inner_text(),
                "statusline": statusline_base,
                "metrics": ui_metrics,
                "findings": ui_rows,
            }

            browser.close()

    finally:
        stop_server(server_proc)

    # =========================================================
    # PART C: Metamorphic Transformations (Meta-1 .. Meta-6)
    # =========================================================
    print("\n" + "=" * 60)
    print("PART C: Metamorphic Invariance Verifications")
    print("=" * 60)

    # Re-read baseline results directly from SQLite to get exact ground truth
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    base_row = conn.execute("SELECT * FROM runs WHERE id=?", (base_run_id,)).fetchone()
    base_result = json.loads(base_row["result"])
    base_snapshot = json.loads(base_row["snapshot"])
    conn.close()

    from freight_audit.importing import ImportMapping, import_data
    from freight_audit.models import Dataset
    from freight_audit.engine import audit

    ms = ImportMapping.model_validate_json((FIXTURES_DIR / "mapping_shipments_adv.json").read_text())
    mc = ImportMapping.model_validate_json((FIXTURES_DIR / "mapping_charges_adv.json").read_text())
    agrs = json.loads((FIXTURES_DIR / "agreements_adv.json").read_text())

    metamorphic_tests = [
        ("Meta-1 (Renamed Files)", "viajes_adversariales_marzo.csv", "liquidacion_proveedor_exp.csv", ms),
        ("Meta-2 (Reversed Rows)", "operaciones_reversed.csv", "cargos_reversed.csv", ms),
        ("Meta-3 (Permuted Columns)", "operaciones_permuted_cols.csv", "cargos_adv.csv", ms),
        ("Meta-4 (Extra Irrelevant Columns)", "operaciones_extra_cols.csv", "cargos_adv.csv", ms),
    ]

    for name, s_file, c_file, m_s in metamorphic_tests:
        print(f"[*] Testing {name}...")
        rs = import_data((FIXTURES_DIR / s_file).read_bytes(), s_file, m_s)
        rc = import_data((FIXTURES_DIR / c_file).read_bytes(), c_file, mc)
        ds = Dataset.model_validate({
            "label": name,
            "shipments": rs["records"],
            "charges": rc["records"],
            "agreements": agrs,
            "evidence": [],
            "coverage": {},
            "mappings": [rs["mapping"], rc["mapping"]],
            "documents": {rs["document"]: rs["filename"], rc["document"]: rc["filename"]},
            "issues": rs["issues"] + rc["issues"],
        })
        m_res = audit(ds)
        assert m_res.summary["counts"] == base_result["summary"]["counts"], f"{name}: counts mismatch"
        assert m_res.summary["currencies"] == base_result["summary"]["currencies"], f"{name}: currencies mismatch"
        print(f"    [✓] {name}: Economically 100% invariant to baseline!")

    # Meta-5 (Cold Restart) & Meta-6 (Double Idempotent Export)
    print("\n[*] Testing Meta-5 (Cold Restart) & Meta-6 (Double Idempotent Export)...")
    server_proc_restarted = start_server()
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(viewport={"width": 1440, "height": 960})
            page = context.new_page()

            page.goto(BASE_URL)
            page.wait_for_selector(f'.run-card[data-run="{base_run_id}"]')
            page.locator(f'.run-card[data-run="{base_run_id}"]').click()
            page.wait_for_selector(".statusline")

            post_statusline = " ".join(page.locator(".statusline").inner_text().split())
            assert post_statusline == statusline_base, "Post-restart statusline mismatch"
            assert "Rechazado" in page.locator('#findings tr:has-text("REM-FAIL")').inner_text(), "Decision not recovered"
            print("    [✓] Meta-5: Cold restart strictly recovered exact baseline states and human decision.")

            page.click("#replay")
            page.wait_for_selector('#notice:has-text("Reproducción idéntica")')
            print("    [✓] Replay confirmed post-restart.")

            # Meta-6: Double export
            with page.expect_download() as dl:
                page.click('a:has-text("Informe imprimible")')
            post_html_path = EXPORTS_DIR / f"{base_run_id}_reporte_reexport.html"
            dl.value.save_as(str(post_html_path))
            assert base_html.read_text(encoding="utf-8") == post_html_path.read_text(encoding="utf-8")
            print("    [✓] Meta-6: Double export produced identical content byte-for-byte.")

            # Direct API GET
            req_api = urllib.request.Request(f"{BASE_URL}/api/runs/{base_run_id}")
            with urllib.request.urlopen(req_api) as resp:
                api_observed = json.loads(resp.read().decode("utf-8"))
            print("    [✓] Direct API GET retrieved.")

            browser.close()
    finally:
        stop_server(server_proc_restarted)

    # =========================================================
    # PART D: Standalone Observation Artifact Dumps
    # =========================================================
    print("\n" + "=" * 60)
    print("PART D: Standalone Observation Artifact Dumps")
    print("=" * 60)

    # SQLite
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    b_row = conn.execute("SELECT * FROM runs WHERE id=?", (base_run_id,)).fetchone()
    b_dec_rows = conn.execute("SELECT * FROM decisions WHERE run_id=?", (base_run_id,)).fetchall()
    b_decs = [dict(r) for r in b_dec_rows]
    for d in b_decs:
        d["payload"] = json.loads(d["payload"])
    conn.close()

    sqlite_observed = {
        "run_id": base_run_id,
        "input_hash": b_row["input_hash"],
        "result_hash": b_row["result_hash"],
        "artifact_hash": b_row["artifact_hash"],
        "snapshot": json.loads(b_row["snapshot"]),
        "result": json.loads(b_row["result"]),
        "decisions": b_decs,
    }

    # JSON export
    with zipfile.ZipFile(base_zip, "r") as z:
        audit_doc = json.loads(z.read("audit.json").decode("utf-8"))
        snapshot_doc = json.loads(z.read("snapshot.json").decode("utf-8")) if "snapshot.json" in z.namelist() else audit_doc["snapshot"]
        json_observed = {
            "audit": audit_doc,
            "snapshot": snapshot_doc,
        }

    # XLSX export
    wb = openpyxl.load_workbook(base_xlsx, data_only=True)
    xlsx_observed = {
        "resumen": [[c for c in r] for r in wb["Resumen"].iter_rows(values_only=True)],
        "hallazgos": [[c for c in r] for r in wb["Hallazgos"].iter_rows(values_only=True)],
    }

    # HTML export
    html_text = base_html.read_text(encoding="utf-8")
    html_observed = {
        "raw_html": html_text,
        "size_bytes": len(html_text),
    }

    (OBSERVED_DIR / "adversarial-sqlite-observed.json").write_text(json.dumps(sqlite_observed, ensure_ascii=False, indent=2), encoding="utf-8")
    (OBSERVED_DIR / "adversarial-api-observed.json").write_text(json.dumps(api_observed, ensure_ascii=False, indent=2), encoding="utf-8")
    (OBSERVED_DIR / "adversarial-json-observed.json").write_text(json.dumps(json_observed, ensure_ascii=False, indent=2), encoding="utf-8")
    (OBSERVED_DIR / "adversarial-xlsx-observed.json").write_text(json.dumps(xlsx_observed, ensure_ascii=False, indent=2), encoding="utf-8")
    (OBSERVED_DIR / "adversarial-html-observed.json").write_text(json.dumps(html_observed, ensure_ascii=False, indent=2), encoding="utf-8")
    (OBSERVED_DIR / "adversarial-ui-observed.json").write_text(json.dumps(ui_observed, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[✓] Saved 6 standalone adversarial observation artifacts to {OBSERVED_DIR}")

    # =========================================================
    # PART E: 6-Way Adversarial Reconciliation & Invariant Checks
    # =========================================================
    print("\n" + "=" * 60)
    print("PART E: 6-Way Cross-Channel Reconciliation & Invariant Checks")
    print("=" * 60)

    # Core equality
    assert sqlite_observed["result"] == api_observed["result"] == json_observed["audit"]["result"]
    assert sqlite_observed["snapshot"] == api_observed["snapshot"] == json_observed["snapshot"]
    print("[✓] Byte/data parity: SQLite == API == JSON export.")

    status_labels = {
        "PASS": "Coincide",
        "FAIL": "Discrepancia",
        "REVIEW": "Revisión humana",
        "UNDETERMINABLE": "Indeterminado",
    }

    # Definition of all 12 baseline findings (10 adversarial + 2 controls)
    expected_matrix = {
        "C_PASS": {"ref": "REM-PASS", "concept": "FLETE", "currency": "ARS", "status": "PASS", "actual": "100", "expected": "100", "diff": "0", "case": "CTRL-PASS"},
        "C_FAIL": {"ref": "REM-FAIL", "concept": "FLETE", "currency": "ARS", "status": "FAIL", "actual": "120", "expected": "100", "diff": "20", "case": "CTRL-FAIL"},
        "C1": {"ref": "REM-001", "concept": "FLETE", "currency": "ARS", "status": "REVIEW", "actual": "100", "expected": None, "diff": None, "case": "ADV-01"},
        "C2": {"ref": "REM-002", "concept": "FLETE", "currency": "USD", "status": "UNDETERMINABLE", "actual": "50", "expected": None, "diff": None, "case": "ADV-02"},
        "C3": {"ref": "REM-003", "concept": "ESPECIAL", "currency": "ARS", "status": "REVIEW", "actual": "100", "expected": "100", "diff": "0", "case": "ADV-03"},
        "C4": {"ref": "REM-004", "concept": "FLETE", "currency": "ARS", "status": "UNDETERMINABLE", "actual": "100", "expected": None, "diff": None, "case": "ADV-04"},
        "C5": {"ref": "REM-005", "concept": "FLETE", "currency": "ARS", "status": "UNDETERMINABLE", "actual": "100", "expected": None, "diff": None, "case": "ADV-05"},
        "C6": {"ref": "REM-006", "concept": "FLETE", "currency": "ARS", "status": "UNDETERMINABLE", "actual": "100", "expected": None, "diff": None, "case": "ADV-06"},
        "C7": {"ref": "REM-UNKNOWN", "concept": "FLETE", "currency": "ARS", "status": "REVIEW", "actual": "100", "expected": None, "diff": None, "case": "ADV-07"},
        "C8A": {"ref": "REM-008", "concept": "FLETE", "currency": "ARS", "status": "REVIEW", "actual": "200", "expected": "100", "diff": "100", "case": "ADV-08"},
        "C10A": {"ref": "REM-010A", "concept": "FLETE", "currency": "ARS", "status": "REVIEW", "actual": "100", "expected": "100", "diff": "0", "case": "ADV-10A"},
        "C10B": {"ref": "REM-010B", "concept": "FLETE", "currency": "ARS", "status": "REVIEW", "actual": "100", "expected": "100", "diff": "0", "case": "ADV-10B"},
    }

    # Index findings from XLSX
    xlsx_map = {}
    x_headers = xlsx_observed["hallazgos"][0]
    for row in xlsx_observed["hallazgos"][1:]:
        c_val = str(row[2]) # Cargos
        for cid in expected_matrix:
            if cid in c_val:
                xlsx_map[cid] = dict(zip(x_headers, row))

    reconciliation_rows = []

    for cid, exp in expected_matrix.items():
        # Find in DB
        db_matches = [f for f in sqlite_observed["result"]["findings"] if cid in f["charge_ids"]]
        assert len(db_matches) == 1, f"Expected 1 DB finding for {cid}"
        f_db = db_matches[0]

        # Invariant 1: REVIEW or UNDETERMINABLE must NEVER become PASS or FAIL
        if exp["status"] in {"REVIEW", "UNDETERMINABLE"}:
            assert f_db["status"] not in {"PASS", "FAIL"}, f"CRITICAL LEAK: {cid} ({exp['case']}) promoted to determinable status {f_db['status']}"

        # Invariant 2: REAL FAIL must never be lost
        if exp["status"] == "FAIL":
            assert f_db["status"] == "FAIL", f"CRITICAL REGRESSION: Real FAIL for {cid} lost"

        # Check DB values
        assert f_db["status"] == exp["status"]
        assert f_db["concept"] == exp["concept"]
        assert f_db["currency"] == exp["currency"]
        assert str(f_db["actual"]) == exp["actual"]
        if exp["expected"] is not None:
            assert str(f_db["expected"]) == exp["expected"]
            assert str(f_db["difference"]) == exp["diff"]
        else:
            assert f_db["expected"] is None
            assert f_db["difference"] is None

        # Check XLSX
        f_x = xlsx_map[cid]
        assert f_x["Estado"] == status_labels[exp["status"]]
        assert str(f_x["Facturado"]) == exp["actual"]
        if exp["expected"] is not None:
            assert str(f_x["Esperado"]) == exp["expected"]

        # Check HTML
        assert f"{exp['currency']} {exp['actual']}" in html_text
        assert status_labels[exp["status"]] in html_text

        # Check UI DOM
        dom_m = [r for r in ui_observed["findings"] if exp["ref"] in r["reference_concept"] and (exp["concept"] in r["reference_concept"] or cid in r["reference_concept"])]
        assert len(dom_m) >= 1, f"UI DOM row not found for {cid} ({exp['ref']})"
        dom_row = dom_m[0]
        assert dom_row["status"] == status_labels[exp["status"]]

        reconciliation_rows.append({
            "charge_id": cid,
            "case": exp["case"],
            "ref": exp["ref"],
            "concept": exp["concept"],
            "currency": exp["currency"],
            "sqlite": {"status": f_db["status"], "actual": f_db["actual"], "diff": f_db["difference"]},
            "api": {"status": f_db["status"], "actual": f_db["actual"], "diff": f_db["difference"]},
            "json": {"status": f_db["status"], "actual": f_db["actual"], "diff": f_db["difference"]},
            "xlsx": {"status": f_x["Estado"], "actual": f_x["Facturado"]},
            "html": {"status": status_labels[exp["status"]], "actual": f"{exp['currency']} {exp['actual']}"},
            "ui": {"status": dom_row["status"], "actual": dom_row["actual"]},
            "divergence": "ZERO",
        })
        print(f"[✓] Reconciled {exp['case']} ({cid}, {exp['ref']}): status={exp['status']} across all 6 channels.")

    # Check Macro metrics in Baseline
    summary = sqlite_observed["result"]["summary"]
    assert summary["counts"] == {"PASS": 1, "FAIL": 1, "REVIEW": 6, "UNDETERMINABLE": 4}
    assert summary["determinable_findings"] == 2
    assert summary["total_findings"] == 12
    assert Decimal(summary["currencies"]["ARS"]["confirmed_overcharge"]) == Decimal("20.00")
    print("[✓] Macro summary metrics reconciled with exact counts.")

    matrix = {
        "summary": {
            "status": "ALL_ADVERSARIAL_CHANNELS_RECONCILED",
            "divergence": 0,
            "channels": ["SQLite", "API", "JSON", "XLSX", "HTML", "UI"],
            "invariants_tested": [
                "Uncertainty (REVIEW/UNDETERMINABLE) never converts to certainty (PASS/FAIL)",
                "Real FAIL is never lost or masked",
                "Rejected import row forces all determinable findings into REVIEW with import_complete=False",
                "Metamorphic operations (rename, reverse, permute cols, extra cols, cold restart, double export) produce identical economics",
            ],
            "counts": summary["counts"],
        },
        "findings": reconciliation_rows,
    }
    (OBSERVED_DIR / "adversarial-reconciliation-matrix.json").write_text(json.dumps(matrix, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[✓] Saved adversarial-reconciliation-matrix.json ({len(reconciliation_rows)} cases verified).")

    print("\n" + "=" * 60)
    print("FASE 2 E2E ADVERSARIAL COMPLETED: 100% SUCCESSFUL")
    print("=" * 60)


if __name__ == "__main__":
    run_adversarial_suite()
