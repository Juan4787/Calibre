"""Read exported representations and compare their material fields with a preserved run."""

import hashlib
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

from .invariants import check_run, read_json, stable_hash

# Deliberately owned by QA, not imported from the renderer being checked.
LABELS = {
    "PASS": "Coincide",
    "FAIL": "Discrepancia",
    "REVIEW": "Revisión humana",
    "UNDETERMINABLE": "Indeterminado",
}
SHEETS = {
    "Resumen",
    "Hallazgos",
    "Cálculos",
    "Problemas de datos",
    "Evidencia",
    "Decisiones",
    "Origen de datos",
    "Metadata",
}


def text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    return str(value)


def rows(values):
    return Counter(tuple(text(cell) for cell in row) for row in values)


class Tables(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables = []
        self.table = None
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in {"td", "th"} and self.cell is not None and self.row is not None:
            self.row.append("".join(self.cell))
            self.cell = None
        elif tag == "tr" and self.row is not None and self.table is not None:
            self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            self.tables.append(self.table)
            self.table = None


def _contained(root: Path, name: str) -> Path:
    path = root / name
    if path.is_symlink() or not path.resolve().is_relative_to(root.resolve()):
        raise ValueError("Export references a path outside its directory")
    return path


def reconcile(directory: Path, *, api_run=None, observed_ui=None) -> dict:
    run = read_json(directory / "audit.json")
    errors = check_run(run)
    channels = {
        "json": "checked",
        "xlsx": "not checked",
        "html": "not checked",
        "api": "not supplied",
        "ui": "not supplied",
    }
    manifest = read_json(directory / "manifest.json")
    if manifest["run_id"] != run["id"] or manifest.get("format") not in {
        "freight-audit-bundle/v1",
        "freight-audit-bundle/v2",
    }:
        errors.append("INV-28 manifest identity mismatch")
    for name, expected in manifest["files"].items():
        path = _contained(directory, name)
        if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != expected:
            errors.append(f"INV-28 export bytes differ: {name}")
    if manifest.get("format") == "freight-audit-bundle/v1":
        snapshot_path = directory / "snapshot.json"
        if not snapshot_path.is_file():
            errors.append("INV-20 missing snapshot.json in v1 bundle")
            snapshot = run["snapshot"]
        else:
            snapshot = read_json(snapshot_path)
            if snapshot != run["snapshot"]:
                errors.append("INV-20 separate snapshot mismatch")
    else:
        if "snapshot.json" in manifest["files"] or (directory / "snapshot.json").exists():
            errors.append("INV-20 redundant snapshot.json in v2 bundle")
        snapshot = run["snapshot"]
    source_hashes = set(snapshot["documents"])
    for entity in ("shipments", "charges"):
        for record in snapshot[entity]:
            source_hashes.update(ref["document"] for ref in record.get("provenance", {}).values())
    source_hashes.update(e["document_hash"] for e in snapshot["evidence"] if e.get("document_hash"))
    for source_hash in source_hashes:
        path = _contained(directory, "sources/" + source_hash)
        if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != source_hash:
            errors.append("INV-22 original source missing or altered")
    findings = run["result"]["findings"]
    workbook_path = directory / "auditoria.xlsx"
    if workbook_path.exists():
        workbook = load_workbook(workbook_path, read_only=True, data_only=False, keep_links=False)
        try:
            if set(workbook.sheetnames) != SHEETS:
                errors.append("INV-23 workbook sheet inventory mismatch")
            elif SHEETS <= set(workbook.sheetnames):

                def compare(sheet, expected, indices=None):
                    actual = []
                    for row in workbook[sheet].iter_rows(min_row=2):
                        if any(cell.data_type == "f" for cell in row):
                            errors.append(f"INV-23 formula in {sheet}")
                        values = [cell.value for cell in row]
                        actual.append(values if indices is None else [values[i] for i in indices])
                    if rows(actual) != rows(expected):
                        errors.append(f"INV-23 XLSX {sheet}: material rows differ")

                compare(
                    "Hallazgos",
                    [
                        [
                            f["id"],
                            ", ".join(f["shipment_ids"]),
                            ", ".join(f["charge_ids"]),
                            f["concept"],
                            f["currency"],
                            f["actual"],
                            f["expected"],
                            f["difference"],
                            f["confirmed_difference"],
                            LABELS[f["status"]],
                            f["agreement"],
                            f["version"],
                            f["rule"],
                            "\n".join(f["reasons"]),
                            ", ".join(f["evidence_ids"]),
                            ", ".join(f["missing_evidence"]),
                        ]
                        for f in findings
                    ],
                )
                summary = run["result"]["summary"]
                wanted = [[c, k, v] for c, bucket in summary["currencies"].items() for k, v in bucket.items()]
                wanted += [["", k, v] for k, v in summary["counts"].items()]
                wanted += [
                    ["", "Importación completa", summary["import_complete"]],
                    ["", "Definición", summary["economic_definition"]],
                ]
                actual_summary = [
                    [cell.value for cell in row] for row in workbook["Resumen"].iter_rows(min_row=2)
                ]
                actual_summary = [row for row in actual_summary if row[1] != "Nota"]
                if rows(actual_summary) != rows(wanted):
                    errors.append("INV-23 XLSX Resumen: metrics differ")
                compare(
                    "Problemas de datos",
                    [
                        [i.get(k) for k in ("category", "document", "row", "field", "raw", "message")]
                        for i in run["result"]["issues"]
                    ],
                )
                compare(
                    "Evidencia",
                    [
                        [
                            e["id"],
                            e["kind"],
                            ", ".join(e["shipment_ids"]),
                            ", ".join(e["charge_ids"]),
                            e["document_hash"],
                            e["note"],
                        ]
                        for e in snapshot["evidence"]
                    ],
                )
                compare(
                    "Decisiones",
                    [
                        [
                            d["seq"],
                            d["payload"]["finding_id"],
                            d["payload"]["action"],
                            d["payload"]["actor"],
                            d["payload"]["note"],
                            d["created_at"],
                            ", ".join(d["payload"]["evidence_ids"]),
                            d["payload"]["known_to_client"],
                            d["payload"]["review_minutes"],
                            d["hash"],
                        ]
                        for d in run["decisions"]
                    ],
                )
                compare(
                    "Origen de datos",
                    [
                        [
                            r["id"],
                            field,
                            ref["filename"],
                            ref["sheet"],
                            ref["row"],
                            ref["column"],
                            ref["raw"],
                            ref["transform"],
                            ref["document"],
                        ]
                        for entity in ("shipments", "charges")
                        for r in snapshot[entity]
                        for field, ref in r["provenance"].items()
                    ],
                )

                def trace_rows(fid, nodes, path=""):
                    for node in nodes:
                        current = path + "/" + node["op"]
                        yield [fid, current, node.get("output"), stable_hash(node.get("details", {}))]
                        yield from trace_rows(fid, node.get("children", []), current)

                from .invariants import json_load

                actual_calcs = [
                    [r[0], r[1], r[3], stable_hash(json_load(r[4]))]
                    for r in workbook["Cálculos"].iter_rows(min_row=2, values_only=True)
                ]
                expected_calcs = [row for f in findings for row in trace_rows(f["id"], f["trace"])]
                if rows(actual_calcs) != rows(expected_calcs):
                    errors.append("INV-23 XLSX Cálculos: trace output/details differ")
                metadata = dict(workbook["Metadata"].iter_rows(min_row=2, values_only=True))
                for key in ("id", "input_hash", "result_hash", "artifact_hash"):
                    if metadata.get(key) != run[key]:
                        errors.append(f"INV-23 XLSX Metadata: {key} mismatch")
                # Economic columns must remain text even if numerically small.
                for row in workbook["Hallazgos"].iter_rows(min_row=2):
                    if any(cell.value is not None and cell.data_type != "s" for cell in row[5:9]):
                        errors.append("INV-14 XLSX money not stored as exact text")
            channels["xlsx"] = "checked"
        finally:
            workbook.close()
    elif (
        "ADVERTENCIA_EXPORTACION.txt" in manifest["files"]
        and (directory / "ADVERTENCIA_EXPORTACION.txt").is_file()
    ):
        channels["xlsx"] = "omitted with explicit warning; JSON available"
    else:
        errors.append("INV-23 XLSX absent without declared warning")
    parser = Tables()
    html = (directory / "reporte.html").read_text(encoding="utf-8")
    parser.feed(html)
    expected_totals = [
        [
            currency,
            *[
                bucket[k]
                for k in (
                    "actual",
                    "determinable",
                    "confirmed_overcharge",
                    "confirmed_undercharge",
                    "review",
                    "undeterminable",
                )
            ],
        ]
        for currency, bucket in run["result"]["summary"]["currencies"].items()
    ]
    expected_findings = [
        [
            ", ".join(f["shipment_ids"]),
            f["concept"],
            LABELS[f["status"]],
            f["currency"] + " " + f["actual"],
            "—" if f["expected"] is None else f["expected"],
            f["confirmed_difference"],
            " ".join(f["reasons"]),
        ]
        for f in findings
    ]
    if (
        len(parser.tables) != 2
        or rows(parser.tables[0][1:]) != rows(expected_totals)
        or rows(parser.tables[1][1:]) != rows(expected_findings)
    ):
        errors.append("INV-23 HTML: visible economic tables differ")
    if run["id"] not in html or run["result_hash"] not in html:
        errors.append("INV-23 HTML: identity missing")
    if not run["result"]["summary"]["import_complete"] and "IMPORTACIÓN INCOMPLETA" not in html:
        errors.append("INV-29 HTML: incomplete import warning missing")
    channels["html"] = "checked"
    if api_run is not None:
        for key in ("id", "input_hash", "result_hash", "artifact_hash", "snapshot", "result", "decisions"):
            if api_run.get(key) != run[key]:
                errors.append(f"INV-23 API: {key} differs (check decision revision)")
        channels["api"] = "checked supplied response"
    if observed_ui is not None:
        # Capture contract, not a claim to have driven a browser. Require all pages.
        expected_ui = {
            "run_id": run["id"],
            "result_hash": run["result_hash"],
            "complete": True,
            "findings": [
                {
                    "id": f["id"],
                    "state": LABELS[f["status"]],
                    "currency": f["currency"],
                    "actual": display_money(f["actual"]),
                    "expected": display_money(f["expected"]),
                    "difference": display_money(f["difference"]),
                }
                for f in findings
            ],
        }
        observed_ui = {**observed_ui, "findings": sorted(observed_ui["findings"], key=lambda f: f["id"])}
        expected_ui["findings"].sort(key=lambda f: f["id"])
        if observed_ui != expected_ui:
            errors.append("INV-23 UI: supplied complete DOM observation differs")
        channels["ui"] = "checked supplied observation; browser capture not authenticated"
    return {
        "run_id": run["id"],
        "errors": errors,
        "channels": channels,
        "limits": "Checks fidelity to preserved result, not truth of agreement. Explanatory prose, CSS and DOM capture authenticity require review.",
    }


def display_money(value):
    if value is None:
        return "—"
    whole, dot, fraction = value.partition(".")
    sign = "-" if whole.startswith("-") else ""
    digits = whole.removeprefix("-")
    chunks: list[str] = []
    while digits:
        chunks.insert(0, digits[-3:])
        digits = digits[:-3]
    return sign + ".".join(chunks) + ("," + fraction if dot else "")
